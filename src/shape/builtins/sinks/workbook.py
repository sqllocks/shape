"""Multi-sheet workbook writer (``[excel]`` extra, openpyxl): one ``.xlsx`` for a whole dataset.

One sheet per table, in the order given (dependency order), then nothing else is guessed: the
sheet names are valid and unique (31 characters, no ``[]:*?/\\``, case-insensitively distinct), and
a ``_README`` sheet, written first so it is the one a reader opens, says what the file is: the
Shape version, seed, domain, schema mode, scale, generation time, the rows of every table, the
table-to-sheet mapping and any chaos or drift planted in the data (read from the chaos
ground-truth log and the drift plan's answer key when given).

Formatting: a styled, frozen header row; column widths fitted to the first rows and capped;
identifier columns (text columns of digits, and text columns that are named like an identifier or
listed in ``identifier_columns``) get the text number format ``@``, so ZIP codes, NDCs, NPIs and
member ids keep their leading zeros when edited in Excel; dates get a date format.

A sheet holds 1,048,576 rows (header included) and 16,384 columns: a larger table is refused,
before anything is written, with an error that names the table. Text that Excel would read as a
formula (``=...``) or an error value (``#N/A``) is stored as text.
"""

from __future__ import annotations

import datetime as dt
import json
import re
from collections.abc import Iterable, Mapping, Sequence
from pathlib import Path
from typing import Any

import pyarrow as pa  # type: ignore[import-untyped]
import pyarrow.compute as pc  # type: ignore[import-untyped]

README_SHEET = "_README"
MAX_SHEET_ROWS = 1_048_576
MAX_SHEET_COLUMNS = 16_384
SHEET_NAME_LIMIT = 31
MAX_COLUMN_WIDTH = 50
MIN_COLUMN_WIDTH = 8
WIDTH_SAMPLE_ROWS = 1_000
BATCH_ROWS = 65_536

_FORBIDDEN_SHEET_CHARS = re.compile(r"[\[\]:*?/\\]")
_IDENTIFIER_NAME = re.compile(
    r"(^|_)(id|ids|key|code|zip|zipcode|postal|postcode|ndc|npi|mrn|ssn|member|account|acct|"
    r"number|num|no)(_|$)|(id|key)$",
    re.IGNORECASE,
)
_DIGITS = re.compile(r"^[0-9]+$")
_SPECIAL_TEXT = re.compile(r"^(=|#(REF!|N/A|VALUE!|DIV/0!|NAME\?|NUM!|NULL!)$)")


class WorkbookTooLargeError(ValueError):
    """A table does not fit on one Excel sheet."""


def valid_sheet_names(
    tables: Sequence[str], reserved: Iterable[str] = (README_SHEET,)
) -> dict[str, str]:
    """``{table: sheet name}``: each name valid for Excel (at most 31 characters, none of
    ``[]:*?/\\``, no leading or trailing apostrophe, not blank), unique ignoring case, and none
    of ``reserved``. A name that has to change keeps its start and gets ``_2``, ``_3`` ... at the
    end."""
    taken = {r.lower() for r in reserved}
    out: dict[str, str] = {}
    for table in tables:
        base = _FORBIDDEN_SHEET_CHARS.sub("_", str(table)).strip().strip("'")
        base = (base or "Sheet")[:SHEET_NAME_LIMIT].rstrip("'")
        name, n = base, 1
        while name.lower() in taken or name.lower() == "history":
            n += 1
            suffix = f"_{n}"
            name = base[: SHEET_NAME_LIMIT - len(suffix)] + suffix
        taken.add(name.lower())
        out[table] = name
    return out


def identifier_columns(table: pa.Table, extra: Iterable[str] = ()) -> list[str]:
    """The columns that get the text number format: text columns whose values are all digits
    (ZIP, NDC without dashes, NPI, member id), text columns named like an identifier, and
    ``extra`` (text columns the schema says are keys)."""
    wanted = set(extra)
    out = []
    for field in table.schema:
        if not (pa.types.is_string(field.type) or pa.types.is_large_string(field.type)):
            continue
        column = table.column(field.name)
        if field.name in wanted or _IDENTIFIER_NAME.search(field.name) or _all_digits(column):
            out.append(field.name)
    return out


def _all_digits(column: pa.ChunkedArray) -> bool:
    non_null = column.drop_null()
    if len(non_null) == 0:
        return False
    return bool(pc.all(pc.match_substring_regex(non_null, _DIGITS.pattern)).as_py())


def _import_openpyxl() -> Any:
    try:
        import openpyxl
    except ImportError as exc:
        raise ImportError(
            "writing Excel needs openpyxl: pip install 'sqllocks-shape[excel]'"
        ) from exc
    return openpyxl


def _plain(value: Any) -> Any:
    """A value Excel can hold: no time zones, no containers or bytes."""
    if isinstance(value, dt.datetime) and value.tzinfo is not None:
        return value.replace(tzinfo=None)
    if isinstance(value, (dict, list, tuple, bytes, set)):
        return str(value)
    return value


class _Writer:
    def __init__(self, workbook: Any) -> None:
        from openpyxl.cell import WriteOnlyCell
        from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
        from openpyxl.styles import Alignment, Font, PatternFill

        self.wb = workbook
        self.cell = WriteOnlyCell
        self.illegal = ILLEGAL_CHARACTERS_RE
        self.header_font = Font(bold=True, color="FFFFFF")
        self.header_fill = PatternFill("solid", start_color="305496", end_color="305496")
        self.header_align = Alignment(vertical="center")
        self.bold = Font(bold=True)

    def text(self, ws: Any, value: str, *, number_format: str | None = None) -> Any:
        value = self.illegal.sub("", value)
        if number_format is None and not _SPECIAL_TEXT.match(value):
            return value
        cell = self.cell(ws, value=value)
        cell.data_type = "s"
        if number_format:
            cell.number_format = number_format
        return cell

    def header(self, ws: Any, names: Sequence[str]) -> list[Any]:
        row = []
        for name in names:
            cell = self.cell(ws, value=self.illegal.sub("", str(name)))
            cell.data_type = "s"
            cell.font = self.header_font
            cell.fill = self.header_fill
            cell.alignment = self.header_align
            row.append(cell)
        return row

    def value(self, ws: Any, value: Any, text_format: bool) -> Any:
        if value is None:
            return None
        value = _plain(value)
        if isinstance(value, str):
            return self.text(ws, value, number_format="@" if text_format else None)
        return value


def _widths(table: pa.Table) -> list[int]:
    head = table.slice(0, WIDTH_SAMPLE_ROWS)
    widths = []
    for field in table.schema:
        longest = len(str(field.name))
        for value in head.column(field.name).to_pylist():
            if value is not None:
                longest = max(longest, len(_width_text(value)))
        widths.append(max(MIN_COLUMN_WIDTH, min(longest + 2, MAX_COLUMN_WIDTH)))
    return widths


def _width_text(value: Any) -> str:
    if isinstance(value, dt.datetime):
        return "2000-01-01 00:00:00"
    if isinstance(value, dt.date):
        return "2000-01-01"
    return str(value)


def check_fits(tables: Mapping[str, pa.Table]) -> None:
    """Raise :class:`WorkbookTooLargeError` for the first table that does not fit on a sheet."""
    for name, table in tables.items():
        if table.num_rows + 1 > MAX_SHEET_ROWS:
            raise WorkbookTooLargeError(
                f"table {name!r} has {table.num_rows:,} rows, and one Excel sheet holds "
                f"{MAX_SHEET_ROWS - 1:,} (plus the header row): write it as csv, parquet or "
                "another format, or generate fewer rows"
            )
        if table.num_columns > MAX_SHEET_COLUMNS:
            raise WorkbookTooLargeError(
                f"table {name!r} has {table.num_columns:,} columns, and one Excel sheet holds "
                f"{MAX_SHEET_COLUMNS:,}"
            )


def _read_chaos_log(source: Any) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    """The ``run`` record and the ``change`` records of a chaos ground-truth log (JSON Lines: a
    ``run`` record, then one ``change`` record per change), from a path or the parsed records."""
    if isinstance(source, (str, Path)):
        records = [
            json.loads(line)
            for line in Path(source).read_text(encoding="utf-8").splitlines()
            if line.strip()
        ]
    else:
        records = list(source)
    run: dict[str, Any] = next((r for r in records if r.get("record") == "run"), {})
    return run, [r for r in records if r.get("record") == "change"]


def _read_drift(source: Any) -> dict[str, Any]:
    """A drift plan or its answer key (``ground_truth.json``), from a path or a dict."""
    if isinstance(source, (str, Path)):
        loaded = json.loads(Path(source).read_text(encoding="utf-8"))
    else:
        loaded = dict(source)
    if not isinstance(loaded, dict):
        raise ValueError("a drift plan or answer key is a JSON object")
    return loaded


def _readme_rows(
    *,
    sheets: Mapping[str, str],
    tables: Mapping[str, pa.Table],
    identifiers: Mapping[str, list[str]],
    meta: Mapping[str, Any],
    chaos: tuple[dict[str, Any], list[dict[str, Any]]] | None,
    drift: dict[str, Any] | None,
) -> list[list[Any]]:
    from shape import __version__

    rows: list[list[Any]] = [["Shape workbook"], []]
    facts = [
        ("Shape version", __version__),
        ("Domain / schema", meta.get("domain") or ""),
        ("Schema mode", meta.get("schema_mode") or ""),
        ("Seed", meta.get("seed")),
        ("Scale", meta.get("scale") or ""),
        ("Generated at (UTC)", meta.get("generated_at")),
        ("Generation time (s)", meta.get("elapsed_seconds")),
        ("Tables", len(tables)),
        ("Total rows", sum(t.num_rows for t in tables.values())),
    ]
    rows += [[k, v] for k, v in facts if v is not None]
    rows += [
        [],
        ["Tables and sheets"],
        ["Table", "Sheet", "Rows", "Columns", "Text-format identifier columns"],
    ]
    for table, sheet in sheets.items():
        t = tables[table]
        rows.append([table, sheet, t.num_rows, t.num_columns, ", ".join(identifiers[table])])
    rows += [
        [],
        ["Identifier columns are stored as text so leading zeros survive; read them as text."],
    ]
    rows += _chaos_rows(chaos)
    rows += _drift_rows(drift)
    if chaos is None and drift is None:
        rows += [[], ["Planted chaos or drift", "none"]]
    return rows


def _chaos_rows(chaos: tuple[dict[str, Any], list[dict[str, Any]]] | None) -> list[list[Any]]:
    if chaos is None:
        return []
    run, changes = chaos
    rows: list[list[Any]] = [[], ["Planted chaos (from the chaos ground-truth log)"]]
    rows.append(["Changes logged", len(changes)])
    for key in ("seed", "batch"):
        if key in run:
            rows.append([f"Chaos {key}", run[key]])
    rows.append(
        ["Table", "Kind", "Column", "Changes", "Example row", "Example before", "Example after"]
    )
    groups: dict[tuple[str, str, str], list[dict[str, Any]]] = {}
    for rec in changes:
        group = (str(rec.get("table")), str(rec.get("kind")), str(rec.get("column") or ""))
        groups.setdefault(group, []).append(rec)
    for (table, kind, column), recs in sorted(groups.items()):
        first = recs[0]
        rows.append(
            [
                table,
                kind,
                column,
                len(recs),
                first.get("row"),
                _show(first.get("before")),
                _show(first.get("after")),
            ]
        )
    return rows


def _drift_rows(drift: dict[str, Any] | None) -> list[list[Any]]:
    if drift is None:
        return []
    rows: list[list[Any]] = [[], ["Planted drift (from the drift plan)"]]
    for key in ("start", "days"):
        if key in drift:
            rows.append([f"Drift {key}", _show(drift[key])])
    events = drift.get("events")
    events = events if isinstance(events, list) else []
    rows.append(["Event", "Kind", "Table.column", "From", "To", "Detail"])
    for i, ev in enumerate(events, 1):
        if not isinstance(ev, dict):
            continue
        target = ev.get("target") or ".".join(
            str(ev[k]) for k in ("table", "column") if ev.get(k) is not None
        )
        detail = {
            k: v
            for k, v in ev.items()
            if k not in ("id", "kind", "table", "column", "target", "start", "end")
        }
        rows.append(
            [
                ev.get("id", i),
                ev.get("kind"),
                target,
                _show(ev.get("start")),
                _show(ev.get("end")),
                json.dumps(detail, sort_keys=True, default=str) if detail else "",
            ]
        )
    return rows


def _show(value: Any) -> Any:
    if value is None or isinstance(value, (int, float, bool)):
        return value
    return str(value)[:200]


def write_workbook(
    path: str | Path,
    tables: Mapping[str, pa.Table],
    *,
    meta: Mapping[str, Any] | None = None,
    identifier_columns_by_table: Mapping[str, Iterable[str]] | None = None,
    chaos_log: Any = None,
    drift_plan: Any = None,
    autofilter: bool = True,
) -> dict[str, str]:
    """Write ``tables`` (name -> Arrow table, in dependency order) as one workbook at ``path``:
    a ``_README`` sheet, then one sheet per table. Returns ``{table: sheet name}``.

    ``meta`` may hold ``domain``, ``schema_mode``, ``seed``, ``scale`` and ``elapsed_seconds`` for
    the ``_README``. ``chaos_log`` is a chaos ground-truth log (path or parsed records) and
    ``drift_plan`` a drift plan or its answer key (path or dict); what they plant is listed in the
    ``_README``. ``identifier_columns_by_table`` adds text-format columns to the detected ones.
    ``autofilter`` (on by default) puts a filter on the header row of every table sheet."""
    openpyxl = _import_openpyxl()
    check_fits(tables)
    info = dict(meta or {})
    info.setdefault("generated_at", dt.datetime.now(dt.UTC).strftime("%Y-%m-%d %H:%M:%S"))
    chaos = _read_chaos_log(chaos_log) if chaos_log is not None else None
    drift = _read_drift(drift_plan) if drift_plan is not None else None
    sheets = valid_sheet_names(list(tables))
    extra = identifier_columns_by_table or {}
    identifiers = {n: identifier_columns(t, extra.get(n, ())) for n, t in tables.items()}
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook(write_only=True)
    try:
        w = _Writer(wb)
        readme = wb.create_sheet(title=README_SHEET)
        for name, table in tables.items():
            _write_sheet(
                w,
                wb.create_sheet(title=sheets[name]),
                table,
                set(identifiers[name]),
                autofilter=autofilter,
            )
        _write_readme(
            w,
            readme,
            _readme_rows(
                sheets=sheets,
                tables=tables,
                identifiers=identifiers,
                meta=info,
                chaos=chaos,
                drift=drift,
            ),
        )
        wb.save(str(target))
    finally:
        wb.close()
    return sheets


def _write_sheet(
    w: _Writer, ws: Any, table: pa.Table, identifiers: set[str], *, autofilter: bool = True
) -> None:
    for i, width in enumerate(_widths(table), 1):
        ws.column_dimensions[_letters(i)].width = width
    ws.freeze_panes = "A2"
    if autofilter and table.num_columns:
        ws.auto_filter.ref = f"A1:{_letters(table.num_columns)}{table.num_rows + 1}"
    ws.append(w.header(ws, table.schema.names))
    text_flags = [f.name in identifiers for f in table.schema]
    for batch in table.to_batches(max_chunksize=BATCH_ROWS):
        columns = [c.to_pylist() for c in batch.columns]
        for record in zip(*columns, strict=True):
            ws.append([w.value(ws, v, flag) for v, flag in zip(record, text_flags, strict=True)])


def _write_readme(w: _Writer, ws: Any, rows: list[list[Any]]) -> None:
    ws.column_dimensions["A"].width = 30
    for letter in "BCDEFG":
        ws.column_dimensions[letter].width = 24
    title_rows = {0}
    section_rows = {
        i
        for i, r in enumerate(rows)
        if len(r) == 1 and i > 0 and r[0] and not str(r[0]).endswith(".")
    }
    for i, row in enumerate(rows):
        cells: list[Any] = []
        for value in row:
            if isinstance(value, str):
                cell = w.text(ws, value)
                if i in title_rows or i in section_rows:
                    cell = w.cell(ws, value=w.illegal.sub("", value))
                    cell.data_type = "s"
                    cell.font = w.bold
                cells.append(cell)
            else:
                cells.append(_plain(value))
        ws.append(cells)


def _letters(index: int) -> str:
    out = ""
    while index:
        index, rem = divmod(index - 1, 26)
        out = chr(65 + rem) + out
    return out
