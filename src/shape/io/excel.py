"""Excel source (``[excel]`` extra, openpyxl in read-only mode): one table per sheet.

``book.xlsx`` is a dataset with one table per visible sheet; ``book.xlsx#Sheet`` is that one sheet
(a hidden sheet is read only when it is named this way, or with ``include_hidden``). The first
non-empty row of a sheet is its header. Rows after the last row that holds a value are not data
(a format alone extends Excel's used range) and are dropped; a sheet that declares a range of
more than 100 million cells is refused with a message that says how to clear it.

Cell types are kept: a cell stored as text stays text (so a ZIP code, an NDC or a member id keeps
its leading zeros), numbers, dates, booleans and blanks come through as their Excel types, and a
column that mixes text with other types becomes text. Cached formula results are read, not the
formulas.

``book.xlsx#Name`` also takes the name of an Excel table or of a named range (a defined name that
points at one block of cells on one sheet); the first row of the block is its header. A sheet of
that name wins over a table, and a table over a named range. Merged cells read as the value of the
top-left cell (the rest are blank) and are reported.

Reading also records what a reader of the workbook would want to know, as *findings* (plain
dicts with ``kind``, ``table``, ``count``, cell positions and examples): numbers and dates stored
as text, hidden columns and sheets, error cells, duplicate or blank headers (renamed
deterministically), merged cells and sentinel values.
"""

from __future__ import annotations

import datetime as dt
import functools
import re
import zipfile
import zlib
from collections import Counter
from collections.abc import Callable
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, ParamSpec, TypeVar

import pyarrow as pa  # type: ignore[import-untyped]

P = ParamSpec("P")
R = TypeVar("R")
XLSX_SUFFIXES = (".xlsx", ".xlsm")
MAX_CELLS = 10  # cell positions kept per finding
MAX_EXAMPLES = 5  # example values kept per finding
ERROR_CODES = ("#REF!", "#N/A", "#VALUE!", "#DIV/0!", "#NAME?", "#NUM!", "#NULL!")
_OLE2 = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"  # legacy .xls, and the wrapper of an encrypted .xlsx

# Placeholder values that stand for "no value" (reported, never removed).
SENTINEL_NUMBERS = (99999, 999999, 9999999, 99999999, -9999, -99999)
SENTINEL_TEXT = {
    "00000": "zero-filled",
    "0000": "zero-filled",
    "000000": "zero-filled",
    "99999": "nines",
    "999999": "nines",
    "9999999": "nines",
    "99999999": "nines",
    "n/a": "not available",
    "na": "not available",
    "null": "null",
    "none": "null",
    "nil": "null",
    "tbd": "to be decided",
    "unknown": "unknown",
    "missing": "missing",
    "-": "dash",
    "--": "dash",
    "?": "question mark",
}
SENTINEL_DATES = (
    dt.date(9999, 12, 31),
    dt.date(1900, 1, 1),
    dt.date(1899, 12, 30),
    dt.date(1, 1, 1),
)

_NUMBER_TEXT = re.compile(
    r"^\s*[-+]?(\d{1,3}(,\d{3})+|\d+)(\.\d+)?([eE][-+]?\d+)?\s*$|^\s*[-+]?\.\d+\s*$"
)
_DATE_TEXT = (
    (re.compile(r"^\d{4}-\d{1,2}-\d{1,2}([ T]\d{1,2}:\d{2}(:\d{2})?)?$"), ("%Y-%m-%d",)),
    (re.compile(r"^\d{4}/\d{1,2}/\d{1,2}$"), ("%Y/%m/%d",)),
    (re.compile(r"^\d{1,2}/\d{1,2}/\d{4}$"), ("%m/%d/%Y", "%d/%m/%Y")),
    (re.compile(r"^\d{1,2}\.\d{1,2}\.\d{4}$"), ("%d.%m.%Y",)),
    (re.compile(r"^\d{1,2}-[A-Za-z]{3}-\d{2,4}$"), ("%d-%b-%Y", "%d-%b-%y")),
    (re.compile(r"^[A-Za-z]{3,9} \d{1,2}, \d{4}$"), ("%B %d, %Y", "%b %d, %Y")),
)


class WorkbookError(ValueError):
    """The file is not a readable ``.xlsx`` workbook (legacy format, password, damaged)."""


# What a workbook whose parts are damaged raises (an XML parser's SyntaxError subclass, ``zlib``
# or ``zipfile`` for a damaged compressed part) before the file is known to be unreadable.
_DAMAGE = (SyntaxError, zlib.error, zipfile.BadZipFile, EOFError)


def _reads_workbook(fn: Callable[P, R]) -> Callable[P, R]:
    """A reader entry point: damage to the workbook's parts is a :class:`WorkbookError` that names
    the file (the first argument)."""

    @functools.wraps(fn)
    def wrapper(*args: P.args, **kwargs: P.kwargs) -> R:
        try:
            return fn(*args, **kwargs)
        except WorkbookError:
            raise
        except _DAMAGE as exc:
            name = Path(str(args[0])).name if args else "the workbook"
            raise WorkbookError(
                f"{name} is damaged ({type(exc).__name__}: {exc}): open it in Excel and save a copy"
            ) from exc

    return wrapper


def is_workbook_path(path: str | Path) -> bool:
    """True for a path whose suffix is ``.xlsx`` or ``.xlsm`` (or the legacy ``.xls`` and
    ``.xlsb``, which are routed here so that they get a clear error)."""
    return Path(path).suffix.lower() in (*XLSX_SUFFIXES, ".xls", ".xlsb")


def split_spec(spec: str | Path) -> tuple[str, str | None]:
    """``"book.xlsx#Sheet"`` -> ``("book.xlsx", "Sheet")``; ``"book.xlsx"`` -> ``("book.xlsx",
    None)``. A ``#`` after anything but a workbook suffix is not a sheet selector."""
    text = str(spec)
    head, sep, tail = text.rpartition("#")
    if sep and tail and is_workbook_path(head):
        return head, tail
    return text, None


def is_workbook_spec(spec: Any) -> bool:
    """True for a string or path that names a workbook, with or without a ``#Sheet`` part."""
    return isinstance(spec, (str, Path)) and is_workbook_path(split_spec(spec)[0])


def _import_openpyxl() -> Any:
    try:
        import openpyxl
    except ImportError as exc:
        raise ImportError(
            "reading Excel needs openpyxl: pip install 'sqllocks-shape[excel]'"
        ) from exc
    return openpyxl


def check_workbook_file(path: str | Path) -> Path:
    """The workbook path, or a :class:`WorkbookError` that says what is wrong with the file."""
    p = Path(path)
    if p.suffix.lower() == ".xls":
        raise WorkbookError(
            f"{p.name} is a legacy .xls workbook, which Shape does not read: open it in Excel "
            "and save it as .xlsx"
        )
    if p.suffix.lower() == ".xlsb":
        raise WorkbookError(
            f"{p.name} is a binary .xlsb workbook, which Shape does not read: save it as .xlsx"
        )
    if not p.is_file():
        raise FileNotFoundError(f"source not found: {p}")
    with open(p, "rb") as fh:
        head = fh.read(8)
    if head == _OLE2:
        raise WorkbookError(
            f"{p.name} is password-protected (or a legacy workbook with an .xlsx name): remove "
            "the password in Excel (File > Info > Protect Workbook) and save it again"
        )
    if not zipfile.is_zipfile(p):
        raise WorkbookError(f"{p.name} is not a valid .xlsx workbook")
    _refuse_zip_bomb(p)
    return p


MAX_EXPANSION = 1000  # an archive member that inflates more than this many times is refused
MAX_PLAIN_BYTES = 256 << 20  # ... once it is also larger than this


def _refuse_zip_bomb(p: Path) -> None:
    try:
        with zipfile.ZipFile(p) as zf:
            infos = zf.infolist()
    except zipfile.BadZipFile as exc:
        raise WorkbookError(f"{p.name} is not a valid .xlsx workbook: {exc}") from exc
    for info in infos:
        if info.file_size > MAX_PLAIN_BYTES and info.file_size > MAX_EXPANSION * max(
            info.compress_size, 1
        ):
            raise WorkbookError(
                f"{p.name} is refused: {info.filename} inflates from {info.compress_size:,} to "
                f"{info.file_size:,} bytes, which is not a spreadsheet"
            )


MAX_SHEET_CELLS = 100_000_000  # a declared range larger than this is refused (rows x columns)


def _refuse_huge_range(ws: Any, sheet: str) -> None:
    """A sheet is read row by row over the range it declares, so a stray or formatted cell at
    the far corner (``A1:XFD1048576`` is 17 billion cells) would run for hours."""
    rows, cols = ws.max_row, ws.max_column
    if rows and cols and rows * cols > MAX_SHEET_CELLS:
        raise WorkbookError(
            f"sheet {sheet!r} declares the range {ws.calculate_dimension()} "
            f"({rows * cols:,} cells), which is too large to read: clear the unused rows and "
            "columns after the data (select them, Clear All) and save a copy"
        )


@dataclass(frozen=True)
class SheetInfo:
    name: str
    state: str  # "visible", "hidden" or "veryHidden"


@_reads_workbook
def sheet_infos(path: str | Path) -> list[SheetInfo]:
    """The sheets of a workbook in workbook order, with their visibility."""
    p = check_workbook_file(path)
    openpyxl = _import_openpyxl()
    wb = _open(openpyxl, p)
    try:
        return [SheetInfo(ws.title, ws.sheet_state) for ws in wb.worksheets]
    finally:
        wb.close()


def _open(openpyxl: Any, p: Path) -> Any:
    try:
        return openpyxl.load_workbook(p, read_only=True, data_only=True)
    except (zipfile.BadZipFile, KeyError, ValueError, OSError) as exc:
        raise WorkbookError(f"{p.name} cannot be read as an .xlsx workbook: {exc}") from exc


@dataclass
class SheetRead:
    """One sheet as an Arrow table, with the findings its cells and header gave."""

    name: str
    table: pa.Table
    findings: list[dict[str, Any]] = field(default_factory=list)


@dataclass
class WorkbookRead:
    """The sheets read from a workbook (in workbook order) and the findings about the workbook
    itself (hidden sheets)."""

    path: str
    sheets: dict[str, SheetRead]
    findings: list[dict[str, Any]] = field(default_factory=list)


def _letters(index: int) -> str:
    """1 -> ``A``, 27 -> ``AA``."""
    out = ""
    while index:
        index, rem = divmod(index - 1, 26)
        out = chr(65 + rem) + out
    return out


def _hidden_columns(path: Path, sheet: str) -> set[int]:
    """1-based indexes of the hidden columns of one sheet (read from the sheet's ``<cols>``)."""
    from openpyxl.xml.functions import iterparse

    try:
        with zipfile.ZipFile(path) as zf:
            target = _sheet_part(zf, sheet)
            if target is None:
                return set()
            hidden: set[int] = set()
            with zf.open(target) as fh:
                for _, el in iterparse(fh, events=("end",)):
                    tag = el.tag.rsplit("}", 1)[-1]
                    if tag == "col" and el.get("hidden") in ("1", "true"):
                        lo, hi = int(el.get("min", "0")), int(el.get("max", "0"))
                        hidden.update(range(lo, min(hi, lo + 16384) + 1))
                    elif tag == "cols" or tag == "sheetData":
                        break
            return hidden
    except (zipfile.BadZipFile, KeyError, ValueError):
        return set()


def _sheet_part(zf: zipfile.ZipFile, sheet: str) -> str | None:
    from openpyxl.xml.functions import fromstring

    book = fromstring(zf.read("xl/workbook.xml"))
    rid = None
    for el in book.iter():
        if el.tag.rsplit("}", 1)[-1] == "sheet" and el.get("name") == sheet:
            rid = next((v for k, v in el.attrib.items() if k.endswith("}id")), None)
    if rid is None:
        return None
    rels = fromstring(zf.read("xl/_rels/workbook.xml.rels"))
    for el in rels:
        if el.get("Id") == rid:
            target = el.get("Target", "")
            return target.lstrip("/") if target.startswith("/") else f"xl/{target}"
    return None


def _is_blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and value == "")


def _unique_headers(raw: list[Any]) -> tuple[list[str], list[dict[str, Any]]]:
    """Header names: text as written, numbers as text, blanks ``column_N``, repeats ``name_2``,
    ``name_3`` ... (bumped past any name already taken). Returns the names and their findings."""
    names: list[str] = []
    notes: list[dict[str, Any]] = []
    seen: set[str] = set()
    base_taken = {str(v) for v in raw if not _is_blank(v)}
    for i, value in enumerate(raw, 1):
        cell = f"{_letters(i)}1"
        if _is_blank(value):
            name = f"column_{i}"
            while name in seen or name in base_taken:
                name += "_"
            notes.append({"kind": "blank_header", "column": name, "cell": cell})
        else:
            name = str(value)
            if name in seen:
                n = 2
                while f"{name}_{n}" in seen or f"{name}_{n}" in base_taken:
                    n += 1
                notes.append(
                    {
                        "kind": "duplicate_header",
                        "column": f"{name}_{n}",
                        "original": name,
                        "cell": cell,
                    }
                )
                name = f"{name}_{n}"
        seen.add(name)
        names.append(name)
    return names, notes


@_reads_workbook
def read_sheet(path: str | Path, sheet: str) -> SheetRead:
    """Read one sheet (hidden or not) into a :class:`SheetRead`."""
    p = check_workbook_file(path)
    openpyxl = _import_openpyxl()
    wb = _open(openpyxl, p)
    try:
        if sheet not in wb.sheetnames:
            raise WorkbookError(f"{p.name} has no sheet {sheet!r}; sheets: {wb.sheetnames}")
        return _read_ws(wb[sheet], p, sheet)
    finally:
        wb.close()


@dataclass(frozen=True)
class Block:
    """A block of cells on one sheet, 1-based and inclusive: an Excel table or a named range."""

    name: str
    kind: str  # "table" or "range"
    sheet: str
    min_col: int
    min_row: int
    max_col: int
    max_row: int


_CELL = re.compile(r"^\$?([A-Za-z]{1,3})\$?(\d+)$")


def _column_index(letters: str) -> int:
    n = 0
    for ch in letters.upper():
        n = n * 26 + ord(ch) - 64
    return n


def _parse_ref(ref: str) -> tuple[int, int, int, int] | None:
    """``"B3:D6"`` (``$`` allowed) -> ``(min_col, min_row, max_col, max_row)``; ``None`` for a
    single cell or anything else."""
    head, sep, tail = ref.partition(":")
    a, b = _CELL.match(head), _CELL.match(tail)
    if not (sep and a and b):
        return None
    c1, r1, c2, r2 = _column_index(a[1]), int(a[2]), _column_index(b[1]), int(b[2])
    return min(c1, c2), min(r1, r2), max(c1, c2), max(r1, r2)


def _split_sheet_ref(text: str) -> tuple[str, str] | None:
    """``"'My Sheet'!$A$1:$B$2"`` -> ``("My Sheet", "$A$1:$B$2")``; ``None`` if there is no
    sheet part."""
    sheet, sep, ref = text.rpartition("!")
    if not sep:
        return None
    if sheet.startswith("'") and sheet.endswith("'") and len(sheet) >= 2:
        sheet = sheet[1:-1].replace("''", "'")
    return sheet, ref


def _read_zip_xml(zf: zipfile.ZipFile, part: str) -> Any:
    from openpyxl.xml.functions import fromstring

    return fromstring(zf.read(part))


def _local(el: Any) -> str:
    return str(el.tag).rsplit("}", 1)[-1]


def _rels_of(zf: zipfile.ZipFile, part: str) -> dict[str, str]:
    """``{relationship id: target part}`` for ``part`` (``xl/worksheets/sheet1.xml``)."""
    folder, _, base = part.rpartition("/")
    rels = f"{folder}/_rels/{base}.rels"
    if rels not in zf.namelist():
        return {}
    out: dict[str, str] = {}
    for el in _read_zip_xml(zf, rels):
        target = el.get("Target", "")
        if target.startswith("/"):
            out[el.get("Id", "")] = target.lstrip("/")
        else:
            parts = [*folder.split("/")]
            for seg in target.split("/"):
                if seg == "..":
                    parts.pop()
                elif seg != ".":
                    parts.append(seg)
            out[el.get("Id", "")] = "/".join(parts)
    return out


@dataclass
class _Names:
    blocks: list[Block] = field(default_factory=list)
    refused: dict[str, str] = field(default_factory=dict)  # name -> why it is not a block


def _find_blocks(path: Path) -> _Names:
    """The Excel tables and the named ranges of a workbook, from its XML (read-only mode of
    openpyxl exposes neither)."""
    names = _Names()
    try:
        with zipfile.ZipFile(path) as zf:
            book = _read_zip_xml(zf, "xl/workbook.xml")
            sheets: list[str] = []
            parts: dict[str, str] = {}
            workbook_rels = _rels_of(zf, "xl/workbook.xml")
            defined: list[Any] = []
            for el in book.iter():
                tag = _local(el)
                if tag == "sheet":
                    title = el.get("name", "")
                    sheets.append(title)
                    rid = next((v for k, v in el.attrib.items() if k.endswith("}id")), None)
                    if rid in workbook_rels:
                        parts[title] = workbook_rels[rid]
                elif tag == "definedName":
                    defined.append(el)
            for title, part in parts.items():
                names.blocks.extend(_table_blocks(zf, title, part, names))
            taken = {b.name.casefold() for b in names.blocks if b.kind == "table"}
            for el in defined:
                _add_range(names, el, sheets, taken)
    except (zipfile.BadZipFile, KeyError, ValueError):
        pass
    return names


def _table_blocks(zf: zipfile.ZipFile, sheet: str, part: str, names: _Names) -> list[Block]:
    if part not in zf.namelist():
        return []
    out: list[Block] = []
    for el in _read_zip_xml(zf, part).iter():
        if _local(el) != "tablePart":
            continue
        rid = next((v for k, v in el.attrib.items() if k.endswith("}id")), None)
        table_part = _rels_of(zf, part).get(rid or "")
        if table_part is None or table_part not in zf.namelist():
            continue
        t = _read_zip_xml(zf, table_part)
        name = t.get("displayName") or t.get("name") or ""
        box = _parse_ref(t.get("ref", ""))
        if not name or box is None:
            continue
        c1, r1, c2, r2 = box
        if t.get("headerRowCount") == "0":
            names.refused[name] = f"table {name!r} has no header row"
            continue
        if t.get("totalsRowCount") not in (None, "0"):
            r2 -= int(t.get("totalsRowCount", "1"))
        out.append(Block(name, "table", sheet, c1, r1, c2, r2))
    return out


def _add_range(names: _Names, el: Any, sheets: list[str], taken: set[str]) -> None:
    name = el.get("name", "")
    if not name or name.startswith("_xlnm.") or el.get("hidden") in ("1", "true"):
        return
    if name.casefold() in taken:
        return
    text = (el.text or "").strip()
    if "," in re.sub(r"'(?:[^']|'')*'", "", text):  # several areas (a comma outside a sheet name)
        names.refused[name] = f"named range {name!r} has more than one area ({text})"
        return
    target = _split_sheet_ref(text)
    if target is None:
        names.refused[name] = f"named range {name!r} is not a block of cells ({text})"
        return
    sheet, ref = target
    if sheet not in sheets:
        names.refused[name] = f"named range {name!r} points at {sheet!r}, which is not a sheet"
        return
    box = _parse_ref(ref)
    if box is None:
        names.refused[name] = f"named range {name!r} is one cell or not a block of cells ({text})"
        return
    names.blocks.append(Block(name, "range", sheet, *box))


@_reads_workbook
def resolve_block(path: Path, name: str, sheets: list[str]) -> Block | None:
    """The table or named range called ``name`` (exact match first, then ignoring case), or
    ``None``. A name that exists but cannot be read as a block raises :class:`WorkbookError`."""
    found = _find_blocks(path)
    for fold in (False, True):

        def same(a: str, b: str, fold: bool = fold) -> bool:
            return a.casefold() == b.casefold() if fold else a == b

        for kind in ("table", "range"):
            for b in found.blocks:
                if b.kind == kind and same(b.name, name):
                    return b
        for refused, why in found.refused.items():
            if same(refused, name):
                raise WorkbookError(f"{path.name}: {why}")
    return None


@_reads_workbook
def read_block(path: str | Path, block: Block) -> SheetRead:
    """Read a table or named range into a :class:`SheetRead` named after it."""
    p = check_workbook_file(path)
    openpyxl = _import_openpyxl()
    wb = _open(openpyxl, p)
    try:
        return _read_ws(wb[block.sheet], p, block.name, block, sheet_name=block.sheet)
    finally:
        wb.close()


@_reads_workbook
def read_selection(path: str | Path, name: str) -> SheetRead:
    """Read what ``book.xlsx#name`` names: a sheet, else an Excel table, else a named range."""
    p = check_workbook_file(path)
    names = [i.name for i in sheet_infos(p)]
    if name in names:
        return read_sheet(p, name)
    block = resolve_block(p, name, names)
    if block is None:
        raise WorkbookError(_no_such(p, name, names))
    return read_block(p, block)


def _no_such(p: Path, name: str, sheets: list[str]) -> str:
    found = _find_blocks(p)
    text = f"{p.name} has no sheet {name!r} (and no table or named range of that name); "
    text += f"sheets: {sheets}"
    if found.blocks:
        text += f"; tables and named ranges: {[b.name for b in found.blocks]}"
    return text


def _merged_ranges(path: Path, sheet: str) -> list[tuple[int, int, int, int]]:
    """The merged ranges of one sheet as ``(min_col, min_row, max_col, max_row)``."""
    from openpyxl.xml.functions import iterparse

    try:
        with zipfile.ZipFile(path) as zf:
            target = _sheet_part(zf, sheet)
            if target is None:
                return []
            out: list[tuple[int, int, int, int]] = []
            with zf.open(target) as fh:
                for _, el in iterparse(fh, events=("end",)):
                    tag = el.tag.rsplit("}", 1)[-1]
                    if tag == "mergeCell":
                        box = _parse_ref(el.get("ref", ""))
                        if box is not None:
                            out.append(box)
                    elif tag == "row":
                        el.clear()
            return out
    except (zipfile.BadZipFile, KeyError, ValueError):
        return []


def _merged_finding(
    path: Path, sheet: str, name: str, area: tuple[int, int, int, int] | None
) -> dict[str, Any] | None:
    """A ``merged_cells`` finding for the merges that touch ``area`` (the whole sheet if
    ``None``): the ranges, their top-left cells (whose value is read) and how many other cells
    they leave blank."""
    ranges: list[tuple[int, int, int, int]] = []
    for c1, r1, c2, r2 in _merged_ranges(path, sheet):
        if area is not None:
            a1, b1, a2, b2 = area
            if c2 < a1 or c1 > a2 or r2 < b1 or r1 > b2:
                continue
        ranges.append((c1, r1, c2, r2))
    if not ranges:
        return None
    ranges.sort(key=lambda b: (b[1], b[0]))

    def label(box: tuple[int, int, int, int]) -> str:
        return f"{_letters(box[0])}{box[1]}:{_letters(box[2])}{box[3]}"

    return {
        "kind": "merged_cells",
        "table": name,
        "count": len(ranges),
        "blank_cells": sum((c2 - c1 + 1) * (r2 - r1 + 1) - 1 for c1, r1, c2, r2 in ranges),
        "ranges": [label(b) for b in ranges[:MAX_CELLS]],
        "cells": [f"{_letters(b[0])}{b[1]}" for b in ranges[:MAX_CELLS]],
    }


def _read_ws(
    ws: Any,
    path: Path,
    sheet: str,
    block: Block | None = None,
    *,
    sheet_name: str | None = None,
) -> SheetRead:
    """One sheet, or ``block`` of it. ``sheet`` names the table; cell positions in findings are
    always positions on the sheet."""
    errors: dict[int, list[tuple[int, str]]] = {}  # column index -> [(row, error code)]
    header: list[Any] | None = None
    header_row = 0
    rows: list[list[Any]] = []
    holds_error: list[bool] = []  # per record: it has an error cell (an error is data)
    if block is None:
        _refuse_huge_range(ws, sheet)
    col0 = block.min_col - 1 if block else 0
    bounds: dict[str, int] = {}
    if block:
        bounds = {
            "min_row": block.min_row,
            "max_row": block.max_row,
            "min_col": block.min_col,
            "max_col": block.max_col,
        }
    for r, row in enumerate(ws.iter_rows(**bounds), block.min_row if block else 1):
        if header is None:
            first = [c.value for c in row]
            if all(_is_blank(v) for v in first):
                continue
            header, header_row = first, r
            while header and _is_blank(header[-1]):
                header.pop()
            continue
        record: list[Any] = []
        has_error = False
        for ci, c in enumerate(row):
            if c.data_type == "e":
                errors.setdefault(ci, []).append((r, str(c.value)))
                record.append(None)
                has_error = True
            else:
                record.append(None if _is_blank(c.value) else c.value)
        rows.append(record)
        holds_error.append(has_error)
    if block is None:
        # cells that only carry a format extend a sheet's used range: blank rows after the last
        # row with a value are not data (an explicit table or named range keeps its rows)
        while rows and not holds_error[-1] and all(v is None for v in rows[-1]):
            rows.pop()
            holds_error.pop()
    width = len(header) if header is not None else 0
    for record in rows:  # data beyond the header's width is kept under a generated name
        while len(record) > width and record[-1] is None:
            record.pop()
        width = max(width, len(record))
    if header is not None:
        header = header + [None] * (width - len(header))
    columns = [[rec[i] if i < len(rec) else None for rec in rows] for i in range(width)]
    findings: list[dict[str, Any]] = []
    if header is None:
        return SheetRead(sheet, pa.table({}), findings)
    names, notes = _unique_headers(header)
    for note in notes:
        cell = note["cell"]
        if col0:  # the header cell on the sheet, not in the block
            note = {**note, "cell": f"{_letters(_column_index(cell[:-1]) + col0)}{header_row}"}
        findings.append({**note, "table": sheet})
    hidden = _hidden_columns(path, sheet_name or sheet)
    for ci, name in enumerate(names):
        if ci + 1 + col0 in hidden:
            findings.append(
                {
                    "kind": "hidden_column",
                    "table": sheet,
                    "column": name,
                    "cell": f"{_letters(ci + 1 + col0)}{header_row}",
                    "column_letter": _letters(ci + 1 + col0),
                }
            )
    arrays = []
    for ci, (name, values) in enumerate(zip(names, columns, strict=True)):
        findings.extend(
            _column_findings(sheet, name, ci + col0, values, header_row, errors.get(ci, []))
        )
        arrays.append(_to_arrow(values))
    area = (block.min_col, block.min_row, block.max_col, block.max_row) if block else None
    merged = _merged_finding(path, sheet_name or sheet, sheet, area)
    if merged is not None:
        findings.append(merged)
    table = pa.table(dict(zip(names, arrays, strict=True))) if names else pa.table({})
    return SheetRead(sheet, table, findings)


def _cell(col: int, row: int) -> str:
    return f"{_letters(col + 1)}{row}"


def _kind_of(value: Any) -> str:
    if isinstance(value, bool):
        return "bool"
    if isinstance(value, int):
        return "int"
    if isinstance(value, float):
        return "float"
    if isinstance(value, str):
        return "str"
    if isinstance(value, dt.datetime):
        return "datetime"
    if isinstance(value, dt.date):
        return "date"
    if isinstance(value, dt.time):
        return "time"
    if isinstance(value, dt.timedelta):
        return "duration"
    return "other"


def _is_number_text(text: str) -> bool:
    return bool(_NUMBER_TEXT.match(text))


def _is_date_text(text: str) -> bool:
    t = text.strip()
    for pattern, formats in _DATE_TEXT:
        if pattern.match(t):
            core = t.split("T")[0].split(" ")[0] if pattern.pattern.startswith("^\\d{4}-") else t
            for fmt in formats:
                try:
                    dt.datetime.strptime(core, fmt)
                except ValueError:
                    continue
                return True
    return False


def _date_of(value: Any) -> dt.date | None:
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    return None


def _finding(
    kind: str,
    sheet: str,
    column: str,
    hits: list[tuple[int, Any]],
    nonblank: int,
    col: int,
    **extra: Any,
) -> dict[str, Any]:
    """``hits`` are ``(row, value)`` pairs in row order."""
    return {
        "kind": kind,
        "table": sheet,
        "column": column,
        "count": len(hits),
        "share": round(len(hits) / nonblank, 6) if nonblank else 0.0,
        "cells": [_cell(col, r) for r, _ in hits[:MAX_CELLS]],
        "examples": [_example(v) for _, v in hits[:MAX_EXAMPLES]],
        **extra,
    }


def _example(value: Any) -> Any:
    if isinstance(value, (dt.datetime, dt.date, dt.time)):
        return value.isoformat()
    if isinstance(value, dt.timedelta):
        return str(value)
    return value


def _column_findings(
    sheet: str,
    name: str,
    ci: int,
    values: list[Any],
    header_row: int,
    errors: list[tuple[int, str]],
) -> list[dict[str, Any]]:
    out: list[dict[str, Any]] = []
    nonblank = sum(v is not None for v in values) + len(errors)
    if errors:
        by_error = Counter(code for _, code in errors)
        out.append(
            {
                "kind": "error_cells",
                "table": sheet,
                "column": name,
                "count": len(errors),
                "share": round(len(errors) / nonblank, 6),
                "cells": [_cell(ci, r) for r, _ in errors[:MAX_CELLS]],
                "examples": [code for _, code in errors[:MAX_EXAMPLES]],
                "by_error": dict(sorted(by_error.items())),
            }
        )
    numbers: list[tuple[int, Any]] = []
    dates: list[tuple[int, Any]] = []
    sentinels: dict[str, list[tuple[int, Any]]] = {}
    kinds: Counter[str] = Counter()
    for i, v in enumerate(values):
        if v is None:
            continue
        row = header_row + 1 + i
        kind = _kind_of(v)
        kinds[kind] += 1
        if kind == "str":
            if _is_number_text(v):
                numbers.append((row, v))
            elif _is_date_text(v):
                dates.append((row, v))
            label = SENTINEL_TEXT.get(v.strip().lower())
            if label is not None:
                sentinels.setdefault(v.strip(), []).append((row, v))
            elif (parsed := _text_sentinel_date(v)) is not None:
                sentinels.setdefault(parsed, []).append((row, v))
        elif kind in ("int", "float") and v in SENTINEL_NUMBERS:
            sentinels.setdefault(str(v), []).append((row, v))
        elif kind in ("datetime", "date") and _date_of(v) in SENTINEL_DATES:
            sentinels.setdefault(str(_date_of(v)), []).append((row, v))
    total = sum(kinds.values()) + len(errors)
    if numbers:
        zeros = sum(
            1
            for _, t in numbers
            if len(t.strip()) > 1 and t.strip()[0] == "0" and not t.strip().startswith("0.")
        )
        out.append(
            _finding(
                "numbers_stored_as_text",
                sheet,
                name,
                numbers,
                total,
                ci,
                leading_zeros=zeros,
                kept_as_text=True,
            )
        )
    if dates:
        out.append(
            _finding("dates_stored_as_text", sheet, name, dates, total, ci, kept_as_text=True)
        )
    for token in sorted(sentinels):
        out.append(
            _finding("sentinel_values", sheet, name, sentinels[token], total, ci, value=token)
        )
    if len(kinds) > 1 and "str" in kinds:
        out.append(
            {
                "kind": "mixed_types",
                "table": sheet,
                "column": name,
                "count": total - kinds["str"],
                "share": round((total - kinds["str"]) / total, 6) if total else 0.0,
                "types": dict(sorted(kinds.items())),
                "stored_as": "string",
            }
        )
    return out


def _text_sentinel_date(text: str) -> str | None:
    t = text.strip()
    for d in SENTINEL_DATES:
        if t.startswith(d.isoformat()):
            return str(d)
    return None


def _to_arrow(values: list[Any]) -> Any:
    kinds = {_kind_of(v) for v in values if v is not None}
    if not kinds:
        return pa.array(values, type=pa.string())
    if kinds == {"str"}:
        return pa.array(values, type=pa.string())
    if kinds == {"bool"}:
        return pa.array(values, type=pa.bool_())
    if kinds == {"int"}:
        if all(v is None or -(2**63) <= v < 2**63 for v in values):
            return pa.array(values, type=pa.int64())
        return pa.array([None if v is None else float(v) for v in values], type=pa.float64())
    if kinds <= {"int", "float"}:
        return pa.array([None if v is None else float(v) for v in values], type=pa.float64())
    if kinds <= {"datetime", "date"}:
        if all(
            v is None or not isinstance(v, dt.datetime) or v.time() == dt.time(0) for v in values
        ):
            return pa.array([_date_of(v) for v in values], type=pa.date32())
        return pa.array([None if v is None else _naive(v) for v in values], type=pa.timestamp("us"))
    if kinds == {"time"}:
        return pa.array(values, type=pa.time64("us"))
    if kinds == {"duration"}:
        return pa.array(values, type=pa.duration("us"))
    return pa.array([None if v is None else _text(v) for v in values], type=pa.string())


def _naive(value: Any) -> dt.datetime:
    if isinstance(value, dt.datetime):
        return value.replace(tzinfo=None)
    return dt.datetime.combine(value, dt.time(0))


def _text(value: Any) -> str:
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, (dt.datetime, dt.date, dt.time)):
        return value.isoformat()
    return str(value)


@_reads_workbook
def read_workbook(
    path: str | Path, sheet: str | None = None, *, include_hidden: bool = False
) -> WorkbookRead:
    """Read a workbook: every visible sheet, or the named ``sheet`` (hidden or not; the name of an
    Excel table or a named range also works, see :func:`read_selection`). Hidden
    sheets are reported, and read only with ``include_hidden`` or when named."""
    p = check_workbook_file(path)
    infos = sheet_infos(p)
    names = [i.name for i in infos]
    block: Block | None = None
    if sheet is not None and sheet not in names:
        block = resolve_block(p, sheet, names)
        if block is None:
            raise WorkbookError(_no_such(p, sheet, names))
    findings = [
        {
            "kind": "hidden_sheet",
            "table": i.name,
            "sheet": i.name,
            "state": i.state,
            "read": sheet == i.name or include_hidden,
        }
        for i in infos
        if i.state != "visible"
    ]
    wanted = (
        [sheet]
        if sheet is not None
        else [i.name for i in infos if i.state == "visible" or include_hidden]
    )
    if not wanted:
        raise WorkbookError(f"{p.name} has no visible sheets")
    if block is not None:
        return WorkbookRead(str(p), {block.name: read_block(p, block)}, findings)
    sheets = {name: read_sheet(p, name) for name in wanted}
    return WorkbookRead(str(p), sheets, findings)


@_reads_workbook
def workbook_sheet_names(path: str | Path, *, include_hidden: bool = False) -> list[str]:
    """The visible sheets (and the hidden ones with ``include_hidden``), in workbook order."""
    return [i.name for i in sheet_infos(path) if i.state == "visible" or include_hidden]
