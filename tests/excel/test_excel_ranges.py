"""Issue #69: named ranges, Excel tables and merged cells as the Excel source; the workbook sink's
autofilter."""

from __future__ import annotations

from pathlib import Path

import pytest

openpyxl = pytest.importorskip("openpyxl")

import pyarrow as pa  # noqa: E402
from openpyxl.workbook.defined_name import DefinedName  # noqa: E402
from openpyxl.worksheet.table import Table  # noqa: E402

import shape  # noqa: E402
from shape.builtins.sinks.workbook import write_workbook  # noqa: E402
from shape.io import WorkbookError, open_source, open_workbook  # noqa: E402
from shape.io.excel import read_workbook  # noqa: E402

ROWS = [
    ["id", "name", "amount"],
    ["00123", "Ada", 10],
    ["00456", "Bo", 20],
    ["00789", "Cy", 30],
]


def _book(path: Path, *, with_noise: bool = True) -> Path:
    """Sheet ``Data`` holds a title row, a block at B3:D6 (a table and a named range) and a note."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws["A1"] = "Quarterly report"
    for r, row in enumerate(ROWS, 3):
        for c, v in enumerate(row, 2):
            ws.cell(row=r, column=c, value=v)
            if isinstance(v, str):
                ws.cell(row=r, column=c).data_type = "s"
    if with_noise:
        ws["B9"] = "footnote"
    ws.add_table(Table(displayName="Members", ref="B3:D6"))
    wb.defined_names["MemberBlock"] = DefinedName("MemberBlock", attr_text="Data!$B$3:$D$6")
    wb.defined_names["Amounts"] = DefinedName("Amounts", attr_text="'Data'!$D$3:$D$6")
    wb.defined_names["TwoAreas"] = DefinedName(
        "TwoAreas", attr_text="Data!$B$3:$B$4,Data!$D$3:$D$4"
    )
    wb.defined_names["Single"] = DefinedName("Single", attr_text="Data!$B$4")
    wb.save(path)
    return path


@pytest.fixture
def book(tmp_path):
    return _book(tmp_path / "book.xlsx")


def test_a_table_is_a_source(book):
    src = open_source(f"{book}#Members")
    table = pa.Table.from_batches(list(src.batches()))
    assert table.column_names == ["id", "name", "amount"]
    assert table.column("id").to_pylist() == ["00123", "00456", "00789"]  # still text
    assert table.column("amount").to_pylist() == [10, 20, 30]


def test_a_named_range_is_a_source_and_its_first_row_is_the_header(book):
    table = pa.Table.from_batches(list(open_source(f"{book}#MemberBlock").batches()))
    assert table.column_names == ["id", "name", "amount"]
    assert table.num_rows == 3


def test_a_one_column_range(book):
    table = pa.Table.from_batches(list(open_source(f"{book}#Amounts").batches()))
    assert table.column_names == ["amount"] and table.num_rows == 3


def test_a_table_name_is_matched_case_insensitively_but_a_sheet_wins(book):
    wb = read_workbook(book, "members")
    assert list(wb.sheets) == ["Members"]
    # a sheet called like a table is the sheet
    w = openpyxl.load_workbook(book)
    w.create_sheet("Members")["A1"] = "x"
    w.save(book)
    assert read_workbook(book, "Members").sheets["Members"].table.column_names == ["x"]


def test_a_range_is_read_by_the_dataset_api_and_the_profile(book):
    assert list(read_workbook(book, "MemberBlock").sheets) == ["MemberBlock"]
    prof = shape.profile(f"{book}#Members")
    assert prof.to_dict()["row_count"] == 3
    assert set(prof.to_dict()["columns"]) == {"id", "name", "amount"}


def test_findings_use_cell_positions_of_the_sheet(tmp_path):
    rows = [["id", "v"], ["1", 5], ["2", "99999"], ["3", 7]]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "S"
    for r, row in enumerate(rows, 4):  # block at C4:D7
        for c, v in enumerate(row, 3):
            ws.cell(row=r, column=c, value=v)
            if isinstance(v, str):
                ws.cell(row=r, column=c).data_type = "s"
    wb.defined_names["Block"] = DefinedName("Block", attr_text="S!$C$4:$D$7")
    path = tmp_path / "f.xlsx"
    wb.save(path)
    sheet = read_workbook(path, "Block").sheets["Block"]
    found = {(f["kind"], f["column"]): f for f in sheet.findings}
    assert found["sentinel_values", "v"]["cells"] == ["D6"]
    assert found["numbers_stored_as_text", "id"]["cells"] == ["C5", "C6", "C7"]


def test_unknown_and_unsupported_names_are_clear_errors(book):
    with pytest.raises(WorkbookError, match="no sheet 'Nope'.*no table or named range"):
        open_source(f"{book}#Nope")
    with pytest.raises(WorkbookError, match="TwoAreas.*more than one area"):
        open_source(f"{book}#TwoAreas")
    with pytest.raises(WorkbookError, match="Single.*one cell"):
        open_source(f"{book}#Single")


def test_a_range_to_a_missing_sheet_is_refused(tmp_path):
    wb = openpyxl.Workbook()
    wb.active["A1"] = "x"
    wb.defined_names["Gone"] = DefinedName("Gone", attr_text="Missing!$A$1:$B$2")
    wb.save(tmp_path / "g.xlsx")
    with pytest.raises(WorkbookError, match="Gone.*Missing"):
        open_source(f"{tmp_path / 'g.xlsx'}#Gone")


def test_a_table_with_a_totals_row_leaves_it_out(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "T"
    for r, row in enumerate([["k", "v"], ["a", 1], ["b", 2], ["Total", 3]], 1):
        for c, v in enumerate(row, 1):
            ws.cell(row=r, column=c, value=v)
    tab = Table(displayName="Sales", ref="A1:B4")
    tab.totalsRowCount = 1
    ws.add_table(tab)
    wb.save(tmp_path / "t.xlsx")
    sheet = read_workbook(tmp_path / "t.xlsx", "Sales").sheets["Sales"]
    assert sheet.table.column("k").to_pylist() == ["a", "b"]


def test_a_table_without_a_header_row_is_refused(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "T"
    ws["A1"], ws["A2"] = 1, 2
    tab = Table(displayName="Bare", ref="A1:A2")
    tab.headerRowCount = 0
    ws.add_table(tab)
    wb.save(tmp_path / "b.xlsx")
    with pytest.raises(WorkbookError, match="Bare.*no header row"):
        open_source(f"{tmp_path / 'b.xlsx'}#Bare")


def test_a_workbook_without_names_still_reads_by_sheet(tmp_path):
    wb = openpyxl.Workbook()
    wb.active.title = "Only"
    wb.active.append(["a"])
    wb.active.append([1])
    wb.save(tmp_path / "p.xlsx")
    assert list(open_workbook(tmp_path / "p.xlsx")) == ["Only"]


# --- merged cells ---------------------------------------------------------------------------------


def _merged(tmp_path) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Orders"
    ws.append(["region", "store", "sales"])
    ws.append(["East", "A", 1])
    ws.append([None, "B", 2])
    ws.append([None, "C", 3])
    ws.append(["West", "D", 4])
    ws.merge_cells("A2:A4")
    wb.save(tmp_path / "m.xlsx")
    return tmp_path / "m.xlsx"


def test_merged_cells_keep_the_top_left_value_and_are_reported(tmp_path):
    sheet = read_workbook(_merged(tmp_path)).sheets["Orders"]
    assert sheet.table.column("region").to_pylist() == ["East", None, None, "West"]
    (f,) = [f for f in sheet.findings if f["kind"] == "merged_cells"]
    assert f["table"] == "Orders"
    assert f["ranges"] == ["A2:A4"]
    assert f["count"] == 1
    assert f["blank_cells"] == 2
    assert f["cells"] == ["A2"]


def test_merged_cells_are_in_the_profile_findings(tmp_path):
    prof = shape.profile(f"{_merged(tmp_path)}#Orders")
    assert [f["kind"] for f in prof.to_dict()["findings"] if f["kind"] == "merged_cells"]


def test_no_merged_finding_without_merges(book):
    sheet = read_workbook(book, "Data").sheets["Data"]
    assert not [f for f in sheet.findings if f["kind"] == "merged_cells"]


def test_only_merges_inside_a_range_are_reported(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "S"
    ws["A1"] = "Title"
    ws.merge_cells("A1:C1")  # outside the block
    for r, row in enumerate([["k", "v"], ["a", 1], [None, 2], ["b", 3]], 3):
        for c, v in enumerate(row, 1):
            ws.cell(row=r, column=c, value=v)
    ws.merge_cells("A4:A5")
    wb.defined_names["Blk"] = DefinedName("Blk", attr_text="S!$A$3:$B$6")
    wb.save(tmp_path / "r.xlsx")
    sheet = read_workbook(tmp_path / "r.xlsx", "Blk").sheets["Blk"]
    (f,) = [f for f in sheet.findings if f["kind"] == "merged_cells"]
    assert f["ranges"] == ["A4:A5"]


# --- autofilter -----------------------------------------------------------------------------------


def _tables() -> dict[str, pa.Table]:
    return {
        "a": pa.table({"x": [1, 2, 3], "y": ["p", "q", "r"]}),
        "b": pa.table({"k": pa.array([], type=pa.int64())}),
    }


def test_autofilter_is_on_by_default(tmp_path):
    path = tmp_path / "f.xlsx"
    write_workbook(path, _tables())
    wb = openpyxl.load_workbook(path)
    assert wb["a"].auto_filter.ref == "A1:B4"
    assert wb["b"].auto_filter.ref == "A1:A1"  # header only
    assert wb["_README"].auto_filter.ref is None


def test_autofilter_can_be_turned_off(tmp_path):
    path = tmp_path / "f.xlsx"
    write_workbook(path, _tables(), autofilter=False)
    wb = openpyxl.load_workbook(path)
    assert wb["a"].auto_filter.ref is None


def test_a_table_with_no_columns_gets_no_autofilter(tmp_path):
    path = tmp_path / "f.xlsx"
    write_workbook(path, {"e": pa.table({})})
    assert openpyxl.load_workbook(path)["e"].auto_filter.ref is None


def test_autofilter_survives_the_round_trip_and_the_generate_option(tmp_path):
    from shape.cli.generation import load_target
    from shape.generation.engine import Engine
    from shape.generation.output import write_engine

    engine = Engine(load_target("retail"), scale="small", seed=7)
    (path,) = write_engine(engine, "excel", tmp_path / "on")
    assert openpyxl.load_workbook(path)["customer"].auto_filter.ref is not None
    engine = Engine(load_target("retail"), scale="small", seed=7)
    (path,) = write_engine(engine, "excel", tmp_path / "off", autofilter=False)
    assert openpyxl.load_workbook(path)["customer"].auto_filter.ref is None


def test_header_findings_point_at_the_sheet_cell_of_a_block(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "S"
    for c, v in enumerate(["k", "k"], 3):
        ws.cell(row=5, column=c, value=v)
    ws.cell(row=6, column=3, value=1)
    wb.defined_names["Blk"] = DefinedName("Blk", attr_text="S!$C$5:$D$6")
    wb.save(tmp_path / "h.xlsx")
    sheet = read_workbook(tmp_path / "h.xlsx", "Blk").sheets["Blk"]
    (f,) = [f for f in sheet.findings if f["kind"] == "duplicate_header"]
    assert f["cell"] == "D5" and f["column"] == "k_2"
