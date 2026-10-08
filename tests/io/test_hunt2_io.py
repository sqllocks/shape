"""Regression tests for the HUNT2-io defects (second audit of the io area)."""

from __future__ import annotations

import datetime as dt
import os
import sys
from pathlib import Path
from typing import Any

import pyarrow as pa  # type: ignore[import-untyped]
import pyarrow.parquet as pq  # type: ignore[import-untyped]
import pytest

from shape.builtins.sinks.files import CsvSink, ParquetSink


def _batches(*values: int) -> list[pa.RecordBatch]:
    return pa.table({"a": list(values)}).to_batches()


def _values(folder: Path) -> list[int]:
    out: list[int] = []
    for name in sorted(os.listdir(folder)):
        out += pq.read_table(folder / name).column("a").to_pylist()
    return out


# ---- #624: append picks the next part from the real part number ---------------------------


def test_append_after_a_table_name_that_contains_00001(tmp_path: Path) -> None:
    for v in (1, 2, 3):
        ParquetSink().write(str(tmp_path), "t00001", _batches(v), roll_rows=1, mode="append")
    assert sorted(os.listdir(tmp_path / "t00001")) == [
        "t00001-00001.parquet",
        "t00001-00002.parquet",
        "t00001-00003.parquet",
    ]
    assert _values(tmp_path / "t00001") == [1, 2, 3]


def test_append_with_a_date_that_contains_00001(tmp_path: Path) -> None:
    template = "{table}/{yyyymmdd}-{part}.{ext}"
    for v in (1, 2, 3):
        ParquetSink().write(
            str(tmp_path),
            "t",
            _batches(v),
            roll_rows=1,
            mode="append",
            path_template=template,
            batch_date="2000-01-01",
        )
    assert _values(tmp_path / "t") == [1, 2, 3]


def test_append_continues_after_part_99999(tmp_path: Path) -> None:
    folder = tmp_path / "t"
    folder.mkdir()
    pq.write_table(pa.table({"a": [1]}), folder / "t-99999.parquet")
    for v in (2, 3):
        ParquetSink().write(str(tmp_path), "t", _batches(v), roll_rows=1, mode="append")
    assert sorted(os.listdir(folder)) == ["t-100000.parquet", "t-100001.parquet", "t-99999.parquet"]


def test_append_ignores_files_that_only_look_like_parts(tmp_path: Path) -> None:
    folder = tmp_path / "t"
    folder.mkdir()
    pq.write_table(pa.table({"a": [1]}), folder / "t-00001.parquet")
    (folder / "t-0000x.parquet").write_bytes(b"")
    (folder / "other-00007.parquet").write_bytes(b"")
    ParquetSink().write(str(tmp_path), "t", _batches(2), roll_rows=1, mode="append")
    assert (folder / "t-00002.parquet").exists()


# ---- #626: Delta commit_rows with an explicit schema ---------------------------------------


def test_delta_commit_rows_with_a_schema_option(tmp_path: Path) -> None:
    pytest.importorskip("deltalake")
    from shape.builtins.sinks.delta import DeltaSink

    schema = pa.schema([("a", pa.int64())])
    rows = DeltaSink().write(
        str(tmp_path),
        "t",
        iter([b for v in (1, 2, 3, 4, 5) for b in _batches(v)]),
        schema=schema,
        commit_rows=2,
    )
    assert rows == 5
    from deltalake import DeltaTable

    table = DeltaTable(str(tmp_path / "t"))
    assert sorted(table.to_pyarrow_table().column("a").to_pylist()) == [1, 2, 3, 4, 5]
    assert table.version() >= 1  # several commits


def test_delta_commit_rows_with_a_schema_and_no_batches(tmp_path: Path) -> None:
    pytest.importorskip("deltalake")
    from shape.builtins.sinks.delta import DeltaSink

    schema = pa.schema([("a", pa.int64())])
    assert DeltaSink().write(str(tmp_path), "t", iter([]), schema=schema, commit_rows=2) == 0
    from deltalake import DeltaTable

    assert DeltaTable(str(tmp_path / "t")).to_pyarrow_table().num_rows == 0


# ---- #625: a plain file is replaced whole or not at all -----------------------------------

_PLAIN = [
    ("csv", "keep.csv"),
    ("tsv", "keep.tsv"),
    ("jsonl", "keep.jsonl"),
    ("parquet", "keep.parquet"),
    ("ipc", "keep.arrow"),
]


def _sink(name: str):  # type: ignore[no-untyped-def]
    from shape.builtins import sinks

    return {
        "csv": sinks.CsvSink,
        "tsv": sinks.TsvSink,
        "jsonl": sinks.JsonlSink,
        "parquet": sinks.ParquetSink,
        "ipc": sinks.IpcSink,
    }[name]()


def _failing(*, schema_change: bool = False):  # type: ignore[no-untyped-def]
    yield pa.table({"a": [1, 2]}).to_batches()[0]
    if schema_change:
        yield pa.table({"a": ["x"]}).to_batches()[0]
    else:
        raise RuntimeError("source failed")


@pytest.mark.parametrize(
    ("fmt", "name", "schema_change"),
    [(f, n, False) for f, n in _PLAIN]
    + [(f, n, True) for f, n in _PLAIN if f in ("parquet", "ipc")],
)  # text formats have no schema to violate
def test_failed_write_keeps_the_previous_file(
    tmp_path: Path, fmt: str, name: str, schema_change: bool
) -> None:
    target = tmp_path / name
    _sink(fmt).write(str(target), "keep", _batches(7))
    before = target.read_bytes()
    with pytest.raises((RuntimeError, pa.ArrowInvalid)):
        _sink(fmt).write(str(target), "keep", _failing(schema_change=schema_change))
    assert target.read_bytes() == before
    assert os.listdir(tmp_path) == [name]  # no temporary file left behind


@pytest.mark.parametrize(("fmt", "name"), _PLAIN)
def test_failed_first_write_leaves_nothing(tmp_path: Path, fmt: str, name: str) -> None:
    target = tmp_path / name
    with pytest.raises(RuntimeError):
        _sink(fmt).write(str(target), "keep", _failing())
    assert os.listdir(tmp_path) == []


@pytest.mark.parametrize(("fmt", "name"), _PLAIN)
def test_successful_write_leaves_only_the_file(tmp_path: Path, fmt: str, name: str) -> None:
    target = tmp_path / name
    assert _sink(fmt).write(str(target), "keep", _batches(1, 2, 3)) == 3
    assert os.listdir(tmp_path) == [name]
    assert _sink(fmt).write(str(target), "keep", _batches(4)) == 1  # replaced, not appended


def test_plain_write_into_a_directory_and_a_long_name(tmp_path: Path) -> None:
    assert CsvSink().write(str(tmp_path) + "/", "t", _batches(1)) == 1
    assert (tmp_path / "t.csv").exists()
    long_name = "x" * 250 + ".csv"
    CsvSink().write(str(tmp_path / long_name), "t", _batches(1))
    assert (tmp_path / long_name).exists()


def test_plain_write_through_a_symlink_and_keeps_the_mode(tmp_path: Path) -> None:
    real = tmp_path / "real.csv"
    real.write_text("old")
    real.chmod(0o640)
    # Windows keeps only the read-only bit of a mode (0o640 reads back as 0o666).
    mode = real.stat().st_mode & 0o777
    assert mode == (0o666 if sys.platform == "win32" else 0o640)
    link = tmp_path / "link.csv"
    link.symlink_to(real)
    CsvSink().write(str(link), "t", _batches(1))
    assert link.is_symlink()
    assert real.read_text() == '"a"\n1\n'
    assert (real.stat().st_mode & 0o777) == mode


@pytest.mark.skipif(not Path("/dev/null").exists(), reason="needs /dev/null")
def test_plain_write_to_a_device_is_written_in_place() -> None:
    assert CsvSink().write("/dev/null", "t", _batches(1, 2)) == 2


# ---- #623: folders that start with _ or . are skipped as a whole ----------------------------


def _tree(root: Path) -> None:
    t = pa.table({"a": [1, 2]})
    for sub in ("", "_shape_tmp", ".hidden", "_delta_log", "keep/deeper"):
        (root / sub).mkdir(parents=True, exist_ok=True)
    pq.write_table(t, root / "p.parquet")
    pq.write_table(t, root / "keep" / "deeper" / "q.parquet")
    pq.write_table(t, root / "_shape_tmp" / "abc-p.parquet")
    pq.write_table(t, root / ".hidden" / "q.parquet")
    pq.write_table(pa.table({"z": [1]}), root / "_delta_log" / "0.checkpoint.parquet")
    pq.write_table(t, root / "_skip.parquet")


def test_directory_read_skips_hidden_and_underscore_folders(tmp_path: Path) -> None:
    from shape.io import expand_paths, read_table

    _tree(tmp_path)
    found = [p.relative_to(tmp_path).as_posix() for p in expand_paths(tmp_path)]
    assert found == ["keep/deeper/q.parquet", "p.parquet"]
    assert read_table(tmp_path).num_rows == 4


def test_a_directory_named_with_an_underscore_can_still_be_the_root(tmp_path: Path) -> None:
    from shape.io import expand_paths

    root = tmp_path / "_landing"
    _tree(root)
    found = [p.relative_to(root).as_posix() for p in expand_paths(root)]
    assert found == ["keep/deeper/q.parquet", "p.parquet"]


def test_abfss_source_skips_hidden_and_underscore_folders() -> None:
    import io as _io

    fsspec = pytest.importorskip("fsspec")
    from shape.builtins.sources.azure import _list_files

    fs = fsspec.filesystem("memory")
    buffer = _io.BytesIO()
    pq.write_table(pa.table({"a": [1]}), buffer)
    for path in (
        "/hunt2/root/part-1.parquet",
        "/hunt2/root/sub/part-2.parquet",
        "/hunt2/root/_shape_tmp/abc-part-3.parquet",
        "/hunt2/root/.hidden/part-4.parquet",
        "/hunt2/root/_delta_log/0.checkpoint.parquet",
    ):
        with fs.open(path, "wb") as handle:
            handle.write(buffer.getvalue())
    try:
        assert _list_files(fs, "hunt2/root") == [
            "/hunt2/root/part-1.parquet",
            "/hunt2/root/sub/part-2.parquet",
        ]
        # a root that is itself below an underscore folder is still read
        assert _list_files(fs, "hunt2/root/_delta_log") == [
            "/hunt2/root/_delta_log/0.checkpoint.parquet"
        ]
    finally:
        fs.rm("/hunt2", recursive=True)


# ---- #627: Delta overwrite replaces the schema too -----------------------------------------


def _delta_rows(path: Path) -> dict[str, list[object]]:
    from deltalake import DeltaTable

    return DeltaTable(str(path)).to_pyarrow_table().sort_by("a").to_pydict()


def test_delta_overwrite_with_a_changed_schema(tmp_path: Path) -> None:
    pytest.importorskip("deltalake")
    from deltalake import DeltaTable

    from shape.builtins.sinks.delta import DeltaSink

    sink = DeltaSink()
    sink.write(str(tmp_path), "t", pa.table({"a": [1, 2]}).to_batches())
    new = pa.table({"a": [3], "b": ["x"]})
    assert sink.write(str(tmp_path), "t", new.to_batches()) == 1
    assert _delta_rows(tmp_path / "t") == {"a": [3], "b": ["x"]}
    protocol = DeltaTable(str(tmp_path / "t")).protocol()
    assert (protocol.min_reader_version, protocol.min_writer_version) == (1, 2)
    assert not protocol.reader_features and not protocol.writer_features
    # and a column removed or retyped
    sink.write(str(tmp_path), "t", pa.table({"a": ["s"]}).to_batches())
    assert DeltaTable(str(tmp_path / "t")).to_pyarrow_table().schema.names == ["a"]


def test_delta_overwrite_in_micro_batches_replaces_the_schema_once(tmp_path: Path) -> None:
    pytest.importorskip("deltalake")
    from shape.builtins.sinks.delta import DeltaSink

    sink = DeltaSink()
    sink.write(str(tmp_path), "t", pa.table({"a": [9]}).to_batches())
    batches = [b for v in (1, 2, 3) for b in pa.table({"a": [v], "b": ["x"]}).to_batches()]
    assert sink.write(str(tmp_path), "t", iter(batches), commit_rows=1) == 3
    assert _delta_rows(tmp_path / "t") == {"a": [1, 2, 3], "b": ["x", "x", "x"]}


def test_delta_append_still_refuses_a_different_schema(tmp_path: Path) -> None:
    pytest.importorskip("deltalake")
    from shape.builtins.sinks.delta import DeltaSink

    sink = DeltaSink()
    sink.write(str(tmp_path), "t", pa.table({"a": [1]}).to_batches())
    with pytest.raises(Exception, match="(?i)schema"):
        sink.write(str(tmp_path), "t", pa.table({"a": [2], "b": ["x"]}).to_batches(), mode="append")
    assert _delta_rows(tmp_path / "t") == {"a": [1]}


# ---- #628: delta+abfss reads the documented AZURE_STORAGE_* variables -------------------------

_DELTA_URI = "delta+abfss://c@acct.dfs.core.windows.net/Tables"


def test_delta_cloud_location_uses_the_storage_key_from_the_environment(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from shape.builtins.sinks.delta import _location

    monkeypatch.setenv("AZURE_STORAGE_ACCOUNT_KEY", "the-key")
    location, storage = _location(_DELTA_URI, "t", {})
    assert location == "abfss://c@acct.dfs.core.windows.net/Tables/t"
    assert storage == {"azure_storage_account_key": "the-key"}


def test_delta_cloud_location_uses_the_sas_token_from_the_environment(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from shape.builtins.sinks.delta import _location

    monkeypatch.delenv("AZURE_STORAGE_ACCOUNT_KEY", raising=False)
    monkeypatch.setenv("AZURE_STORAGE_SAS_TOKEN", "sig=1")
    assert _location(_DELTA_URI, "t", {})[1] == {"azure_storage_sas_key": "sig=1"}


def test_delta_cloud_location_prefers_an_explicit_credential_to_the_environment(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from shape.builtins.sinks.delta import _location

    monkeypatch.setenv("AZURE_STORAGE_ACCOUNT_KEY", "from-env")
    assert _location(_DELTA_URI, "t", {"account_key": "explicit"})[1] == {
        "azure_storage_account_key": "explicit"
    }
    storage = _location(_DELTA_URI, "t", {"token": "tok"})[1]
    assert storage == {"azure_storage_token": "tok"}


def test_delta_cloud_location_says_a_connection_string_cannot_open_a_delta_table(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from shape.builtins.sinks.delta import _location

    for name in ("AZURE_STORAGE_ACCOUNT_KEY", "AZURE_STORAGE_SAS_TOKEN"):
        monkeypatch.delenv(name, raising=False)
    monkeypatch.setenv("AZURE_STORAGE_CONNECTION_STRING", "UseDevelopmentStorage=true")
    with pytest.raises(ValueError, match="connection string"):
        _location(_DELTA_URI, "t", {})


# ---- #629: the single-table Excel sink stores text as text ----------------------------------


def _sheet_cells(path: Path) -> tuple[str, list[list[tuple[object, str]]]]:
    openpyxl = pytest.importorskip("openpyxl")
    wb = openpyxl.load_workbook(path)
    ws = wb.worksheets[0]
    return ws.title, [[(c.value, c.data_type) for c in row] for row in ws.iter_rows()]


def test_excel_sink_stores_formula_like_text_as_text(tmp_path: Path) -> None:
    from shape.builtins.sinks.excel import ExcelSink

    t = pa.table({"a": ["=1+1", "#N/A", "@SUM(1)", "+1", "ok"], "=h": [1, 2, 3, 4, 5]})
    ExcelSink().write(str(tmp_path / "x.xlsx"), "x", t.to_batches())
    _, rows = _sheet_cells(tmp_path / "x.xlsx")
    assert rows[0] == [("a", "s"), ("=h", "s")]
    assert [r[0] for r in rows[1:]] == [
        ("=1+1", "s"),
        ("#N/A", "s"),
        ("@SUM(1)", "s"),
        ("+1", "s"),
        ("ok", "s"),
    ]
    assert [r[1] for r in rows[1:]] == [(1, "n"), (2, "n"), (3, "n"), (4, "n"), (5, "n")]


def test_excel_sink_removes_characters_a_worksheet_cannot_hold(tmp_path: Path) -> None:
    from shape.builtins.sinks.excel import ExcelSink

    t = pa.table({"a": ["a\x01b", "tab\there", "é\u2028x"]})
    assert ExcelSink().write(str(tmp_path / "x.xlsx"), "x", t.to_batches()) == 3
    _, rows = _sheet_cells(tmp_path / "x.xlsx")
    assert [r[0][0] for r in rows[1:]] == ["ab", "tab\there", "é\u2028x"]


@pytest.mark.parametrize(
    ("table", "sheet"),
    [
        ("a/b", "a_b"),
        ("a:b*c?", "a_b_c_"),
        ("[x]", "_x_"),
        ("'quoted'", "quoted"),
        ("x" * 40, "x" * 31),
        ("", "Sheet"),
    ],
)
def test_excel_sink_makes_the_table_name_a_valid_sheet_name(
    tmp_path: Path, table: str, sheet: str
) -> None:
    from shape.builtins.sinks.excel import ExcelSink

    target = tmp_path / "x.xlsx"
    ExcelSink().write(str(target), table, pa.table({"a": [1]}).to_batches())
    assert _sheet_cells(target)[0] == sheet


def test_excel_sink_keeps_a_plain_table_name_and_empty_tables(tmp_path: Path) -> None:
    from shape.builtins.sinks.excel import ExcelSink

    target = tmp_path / "x.xlsx"
    assert (
        ExcelSink().write(str(target), "orders", iter([]), schema=pa.schema([("a", pa.int64())]))
        == 0
    )
    title, rows = _sheet_cells(target)
    assert title == "orders"
    assert rows == [[("a", "s")]]


# ---- #630: JSON Lines output is valid JSON --------------------------------------------------


def _strict_lines(path: Path) -> list[dict[str, object]]:
    import json

    def refuse(constant: str) -> None:
        raise ValueError(f"{constant} is not JSON")

    return [
        json.loads(line, parse_constant=refuse) for line in path.read_text("utf-8").splitlines()
    ]


def test_jsonl_writes_non_finite_floats_as_null(tmp_path: Path) -> None:
    from shape.builtins.sinks.files import JsonlSink

    t = pa.table(
        {
            "f": [float("nan"), float("inf"), float("-inf"), 1.5, -0.0, 1e300, None],
            "n": [[1.0, float("nan")], None, [], [float("inf")], [2.5], [0.0], None],
            "s": [{"x": float("nan"), "y": 1.0}, None, None, None, None, None, None],
        }
    )
    JsonlSink().write(str(tmp_path / "o.jsonl"), "o", t.to_batches())
    rows = _strict_lines(tmp_path / "o.jsonl")
    assert [r["f"] for r in rows] == [None, None, None, 1.5, -0.0, 1e300, None]
    assert [r["n"] for r in rows] == [[1.0, None], None, [], [None], [2.5], [0.0], None]
    assert rows[0]["s"] == {"x": None, "y": 1.0}


def test_jsonl_writes_binary_as_base64(tmp_path: Path) -> None:
    import base64

    from shape.builtins.sinks.files import JsonlSink

    t = pa.table({"b": pa.array([b"\x00\xff", b"", None, b"abc"], pa.binary())})
    JsonlSink().write(str(tmp_path / "o.jsonl"), "o", t.to_batches())
    values = [r["b"] for r in _strict_lines(tmp_path / "o.jsonl")]
    assert values == ["AP8=", "", None, "YWJj"]
    assert base64.b64decode(str(values[0])) == b"\x00\xff"


def test_jsonl_keeps_the_other_values_as_they_were(tmp_path: Path) -> None:
    import decimal

    from shape.builtins.sinks.files import JsonlSink

    t = pa.table(
        {
            "s": ["é😀", None],
            "i": [2**62, -1],
            "d": pa.array([decimal.Decimal("1.50"), None], pa.decimal128(10, 2)),
            "t": pa.array([dt.datetime(2026, 1, 2, 3, 4, 5), None]),
            "dd": pa.array([dt.date(2026, 1, 2), None]),
        }
    )
    JsonlSink().write(str(tmp_path / "o.jsonl"), "o", t.to_batches())
    assert (tmp_path / "o.jsonl").read_text("utf-8").splitlines() == [
        '{"s":"é😀","i":4611686018427387904,"d":"1.50","t":"2026-01-02T03:04:05","dd":"2026-01-02"}',
        '{"s":null,"i":-1,"d":null,"t":null,"dd":null}',
    ]


# ---- #719, #720: the Excel reader and unused cells --------------------------------------------


def _write_book(path: Path, build) -> Path:  # type: ignore[no-untyped-def]
    openpyxl = pytest.importorskip("openpyxl")
    wb = openpyxl.Workbook()
    build(wb.active, openpyxl)
    wb.save(path)
    return path


def _styled(ws, openpyxl, *cells: str) -> None:  # type: ignore[no-untyped-def]
    for ref in cells:
        ws[ref].fill = openpyxl.styles.PatternFill("solid", start_color="FFFF00")


def test_excel_reader_drops_formatted_but_empty_trailing_rows(tmp_path: Path) -> None:
    from shape.io import read_table

    def build(ws, openpyxl):  # type: ignore[no-untyped-def]
        ws.append(["h1", "h2"])
        ws.append([1, "a"])
        ws.append([None, None])  # a blank row between data rows is kept
        ws.append([2, "b"])
        _styled(ws, openpyxl, *(f"A{r}" for r in range(10, 16)))

    table = read_table(_write_book(tmp_path / "t.xlsx", build))
    assert table.to_pydict() == {"h1": [1, None, 2], "h2": ["a", None, "b"]}


def test_excel_reader_header_only_sheet_with_formatted_rows_has_no_rows(tmp_path: Path) -> None:
    from shape.io import read_table

    def build(ws, openpyxl):  # type: ignore[no-untyped-def]
        ws.append(["h1", "h2"])
        _styled(ws, openpyxl, "A5", "B6")

    table = read_table(_write_book(tmp_path / "t.xlsx", build))
    assert table.num_rows == 0
    assert table.column_names == ["h1", "h2"]


def test_excel_reader_keeps_trailing_rows_with_a_value_or_an_error(tmp_path: Path) -> None:
    from shape.io import read_table

    def build(ws, openpyxl):  # type: ignore[no-untyped-def]
        ws.append(["h1", "h2"])
        ws.append([1, "a"])
        ws.append([None, "#N/A"])  # an error cell is data
        ws["B3"].data_type = "e"
        ws.append([None, " "])  # whitespace is a value
        _styled(ws, openpyxl, "A9")

    table = read_table(_write_book(tmp_path / "t.xlsx", build))
    assert table.num_rows == 3
    assert table.column("h2").to_pylist() == ["a", None, " "]


def test_excel_reader_keeps_the_blank_rows_of_an_explicit_named_range(tmp_path: Path) -> None:
    from shape.io import read_table

    def build(ws, openpyxl):  # type: ignore[no-untyped-def]
        from openpyxl.workbook.defined_name import DefinedName

        ws.title = "Data"
        ws.append(["h1", "h2"])
        ws.append([1, "a"])
        _styled(ws, openpyxl, "A4", "A5")  # the rows exist, with a format and no value
        wb = ws.parent
        wb.defined_names["block"] = DefinedName("block", attr_text="Data!$A$1:$B$5")

    table = read_table(str(_write_book(tmp_path / "t.xlsx", build)) + "#block")
    assert table.num_rows == 4


class _TooSlow(Exception):
    pass


def _within(seconds: int):  # type: ignore[no-untyped-def]
    import contextlib
    import signal

    @contextlib.contextmanager
    def guard():  # type: ignore[no-untyped-def]
        if not hasattr(signal, "SIGALRM"):
            yield
            return

        def fire(*_: object) -> None:
            raise _TooSlow(f"still running after {seconds} s")

        old = signal.signal(signal.SIGALRM, fire)
        signal.alarm(seconds)
        try:
            yield
        finally:
            signal.alarm(0)
            signal.signal(signal.SIGALRM, old)

    return guard()


def test_excel_reader_refuses_a_sheet_with_a_huge_declared_range(tmp_path: Path) -> None:
    from shape.io import read_table
    from shape.io.excel import WorkbookError

    def build(ws, openpyxl):  # type: ignore[no-untyped-def]
        ws.append(["h"])
        ws.append([1])
        ws["XFD1048576"] = "far"

    path = _write_book(tmp_path / "far.xlsx", build)
    with _within(30), pytest.raises(WorkbookError, match=r"A1:XFD1048576"):
        read_table(path)


def test_excel_reader_still_reads_a_wide_and_a_tall_sheet(tmp_path: Path) -> None:
    from shape.io import read_table

    def build(ws, openpyxl):  # type: ignore[no-untyped-def]
        ws.append([f"c{i}" for i in range(200)])
        for r in range(2000):
            ws.append([r] * 200)

    with _within(60):
        table = read_table(_write_book(tmp_path / "big.xlsx", build))
    assert (table.num_rows, table.num_columns) == (2000, 200)


# ---- #724, #725: the SQL sink ----------------------------------------------------------------


def _script(tmp_path: Path, table: pa.Table, **options: object) -> str:
    from shape.builtins.sinks.sql import SqlSink

    SqlSink().write(str(tmp_path / "o.sql"), "t", table.to_batches(), **options)
    return (tmp_path / "o.sql").read_bytes().decode("utf-8")


def _go_lines(script: str) -> int:
    import re

    return len(re.findall(r"(?im)(?:^|[\r\n])[ \t]*go(?:\W|$)", script))


def test_tsql_value_with_a_go_line_does_not_split_the_batch(tmp_path: Path) -> None:
    t = pa.table({"s": ["x\nGO\nDROP TABLE victim;\nGO\nSELECT '"]})
    script = _script(tmp_path, t, ddl=False)
    assert _go_lines(script) == 1  # the separator the sink writes after the INSERT
    expected = (
        "(N'x' + NCHAR(10) + N'GO' + NCHAR(10) + N'DROP TABLE victim;' + NCHAR(10) + N'GO'"
        " + NCHAR(10) + N'SELECT ''')"
    )
    assert expected in script


def test_fabric_warehouse_value_with_a_go_line_uses_char(tmp_path: Path) -> None:
    t = pa.table({"s": ["a\r\n  go 3\r\nb"]})
    script = _script(tmp_path, t, ddl=False, sql_dialect="tsql-fabric-warehouse")
    assert _go_lines(script) == 1
    assert "('a' + CHAR(13) + CHAR(10) + '  go 3' + CHAR(13) + CHAR(10) + 'b')" in script


@pytest.mark.parametrize("dialect", ["postgres", "mysql"])
def test_other_dialects_keep_line_breaks_in_a_literal(tmp_path: Path, dialect: str) -> None:
    t = pa.table({"s": ["x\nGO\ny"]})
    assert "('x\nGO\ny')" in _script(tmp_path, t, ddl=False, sql_dialect=dialect)


def test_tsql_text_without_a_go_line_is_unchanged(tmp_path: Path) -> None:
    t = pa.table({"s": ["a\nb", "GOTO\nlabel", "ago\nx", "-- go", "it's", "x\ngo home"]})
    script = _script(tmp_path, t, ddl=False)
    assert "(N'a\nb')" in script
    assert "(N'GOTO\nlabel')" in script  # GOTO is not GO
    assert "(N'x' + NCHAR(10) + N'go home')" in script  # a line that starts with go is cut
    assert "(N'ago\nx')" in script
    assert "(N'-- go')" in script
    assert "(N'it''s')" in script
    assert _go_lines(script) == 1


@pytest.mark.parametrize("name", ["a\nGO\nb", "a\rb"])
def test_tsql_refuses_a_name_with_a_line_break(tmp_path: Path, name: str) -> None:
    from shape.builtins.sinks.sql import SqlSink

    with pytest.raises(ValueError, match="line break"):
        SqlSink().write(str(tmp_path / "o.sql"), "t", pa.table({name: [1]}).to_batches(), ddl=False)
    with pytest.raises(ValueError, match="line break"):
        SqlSink().write(str(tmp_path / "o.sql"), name, pa.table({"a": [1]}).to_batches())
    # other dialects can quote such a name
    assert name in _script(tmp_path, pa.table({name: [1]}), ddl=False, sql_dialect="postgres")


@pytest.mark.parametrize(
    ("key", "message"),
    [(["zz"], r"\['zz'\].*not columns of table 't'"), (["k", "k"], "more than once"), ("zz", "zz")],
)
def test_primary_key_must_name_columns_of_the_table(
    tmp_path: Path, key: object, message: str
) -> None:
    from shape.builtins.sinks.sql import SqlSink

    with pytest.raises(ValueError, match=message):
        SqlSink().write(
            str(tmp_path / "o.sql"), "t", pa.table({"k": [1]}).to_batches(), primary_key=key
        )


def test_primary_key_as_one_string_is_one_column(tmp_path: Path) -> None:
    script = _script(tmp_path, pa.table({"key": [1]}), primary_key="key")
    assert "PRIMARY KEY ([key])" in script


# ---- #625 (continued): the script and workbook sinks replace their file whole -----------------


def test_sql_and_excel_sinks_keep_the_previous_file_when_a_write_fails(tmp_path: Path) -> None:
    pytest.importorskip("openpyxl")
    from shape.builtins.sinks.excel import ExcelSink
    from shape.builtins.sinks.sql import SqlSink

    for sink, name in ((SqlSink(), "keep.sql"), (ExcelSink(), "keep.xlsx")):
        target = tmp_path / name
        sink.write(str(target), "keep", _batches(7))
        before = target.read_bytes()
        with pytest.raises(RuntimeError):
            sink.write(str(target), "keep", _failing())
        assert target.read_bytes() == before
    assert sorted(os.listdir(tmp_path)) == ["keep.sql", "keep.xlsx"]  # no temporary files


def test_sql_sink_refusing_a_primary_key_leaves_no_file(tmp_path: Path) -> None:
    from shape.builtins.sinks.sql import SqlSink

    with pytest.raises(ValueError, match="primary_key"):
        SqlSink().write(
            str(tmp_path / "o.sql"), "t", pa.table({"k": [1]}).to_batches(), primary_key=["zz"]
        )
    assert os.listdir(tmp_path) == []


# ---- #734: duplicate or empty CSV column names ------------------------------------------------


@pytest.mark.parametrize(
    ("text", "names"),
    [
        ("a,a,a\n1,2,3\n", "['a']"),
        ("a,a\nx,y\n", "['a']"),
        (",,\n1,2,3\n", "['']"),
        ("a,b,a,b\n1,2,3,4\n", "['a', 'b']"),
    ],
)
@pytest.mark.parametrize("stream", [False, True])
def test_csv_with_duplicate_names_is_refused_clearly(
    tmp_path: Path, text: str, names: str, stream: bool
) -> None:
    from shape.io import CsvOptions, ReaderError, open_source, read_table

    path = tmp_path / "dup.csv"
    path.write_text(text)
    with pytest.raises(ReaderError, match=r"duplicate column names") as caught:
        read_table(path, csv=CsvOptions(stream=stream))
    assert names in str(caught.value)
    assert "dup.csv" in str(caught.value)
    with pytest.raises(ReaderError, match="duplicate column names"):
        open_source(path, csv=CsvOptions(stream=stream)).table()


def test_csv_with_unique_names_and_a_named_header_still_reads(tmp_path: Path) -> None:
    from shape.io import CsvOptions, read_table

    path = tmp_path / "ok.csv"
    path.write_text("a,b\n1,x\n")
    assert read_table(path).to_pydict() == {"a": [1], "b": ["x"]}
    # headerless files name their own columns
    nohead = tmp_path / "nohead.csv"
    nohead.write_text("1,x\n2,y\n")
    table = read_table(nohead, csv=CsvOptions(has_header=False))
    assert table.num_rows == 2
    named = read_table(nohead, csv=CsvOptions(has_header=False, column_names=("k", "v")))
    assert named.column_names == ["k", "v"]


def test_profile_of_a_csv_with_duplicate_names_keeps_every_column(tmp_path: Path) -> None:
    # #734 (no column dropped without a word) as int/INT-19 settled it for the profiler: #167
    # names repeated and blank header names as pandas does, so every column is kept. The CSV
    # readers of shape.io still refuse a repeated name (above).
    import shape

    for text, names in (
        ("a,a\nx,y\n", ["a", "a.1"]),
        ("a,a,a\n1,2,3\n", ["a", "a.1", "a.2"]),
        (",,\n1,2,3\n", ["Unnamed: 0", "Unnamed: 1", "Unnamed: 2"]),
    ):
        path = tmp_path / "dup.csv"
        path.write_text(text)
        assert list(shape.profile(str(path)).to_dict()["columns"]) == names


# ---- #736: the fabric-mirror sink ------------------------------------------------------------


@pytest.mark.parametrize("uri", ["postgresql://host/db", "s3://bucket/x", "https://h/c"])
def test_fabric_mirror_refuses_a_scheme_it_does_not_write(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, uri: str
) -> None:
    from shape.builtins.sinks.fabric_mirror import FabricMirrorSink
    from shape.plugins.schemes import UnsupportedSchemeError

    monkeypatch.chdir(tmp_path)
    with pytest.raises(UnsupportedSchemeError, match="fabric-mirror sink"):
        FabricMirrorSink().write(uri, "t", _batches(1))
    assert os.listdir(tmp_path) == []


def test_fabric_mirror_still_writes_a_path_and_a_file_uri(tmp_path: Path) -> None:
    from shape.builtins.sinks.fabric_mirror import FabricMirrorSink

    assert FabricMirrorSink().write(str(tmp_path / "a"), "t", _batches(1)) == 1
    assert FabricMirrorSink().write((tmp_path / "b").as_uri(), "t", _batches(1, 2)) == 2
    assert os.listdir(tmp_path / "b" / "t") == ["00000000000000000001.parquet"]


@pytest.mark.parametrize("damaged", ["{not json", "", "[1]", "null", '{"keyColumns": 5}'])
def test_fabric_mirror_names_a_damaged_metadata_file(tmp_path: Path, damaged: str) -> None:
    from shape.builtins.sinks.fabric_mirror import FabricMirrorSink

    sink = FabricMirrorSink()
    table = pa.table({"k": [1], "v": ["a"]})
    sink.write(str(tmp_path), "t", table.to_batches(), key_columns=["k"])
    (tmp_path / "t" / "_metadata.json").write_text(damaged)
    with pytest.raises(ValueError, match="_metadata.json"):
        sink.write(str(tmp_path), "t", table.to_batches(), key_columns=["k"])
    assert sorted(os.listdir(tmp_path / "t")) == ["00000000000000000001.parquet", "_metadata.json"]


def test_fabric_mirror_publish_moves_the_file(tmp_path: Path) -> None:
    from shape.builtins.sinks.fabric_mirror import _Local

    final = tmp_path / "00000000000000000001.parquet"
    temp = tmp_path / "_00000000000000000001.parquet"
    temp.write_bytes(b"mine")
    _Local(tmp_path).publish(str(temp), str(final))
    assert final.read_bytes() == b"mine"
    assert not temp.exists()


# ---- #738: every sink entry point checks the scheme -------------------------------------------


@pytest.mark.parametrize("uri", ["postgresql://host/db", "s3://bucket/x"])
def test_open_table_and_write_workbook_refuse_other_schemes(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, uri: str
) -> None:
    from shape.builtins import sinks
    from shape.plugins.schemes import UnsupportedSchemeError

    monkeypatch.chdir(tmp_path)
    for sink in (sinks.CsvSink(), sinks.TsvSink(), sinks.JsonlSink(), sinks.ParquetSink()):
        with pytest.raises(UnsupportedSchemeError, match=f"the {sink.name} sink"):
            sink.open_table(uri, "t", roll_rows=1)
    pytest.importorskip("openpyxl")
    with pytest.raises(UnsupportedSchemeError, match="the excel sink"):
        sinks.ExcelSink().write_workbook(uri + "/x.xlsx", {"t": pa.table({"a": [1]})})
    assert os.listdir(tmp_path) == []


def test_open_table_still_takes_a_path_and_a_file_uri(tmp_path: Path) -> None:
    from shape.builtins.sinks.files import ParquetSink

    for uri in (str(tmp_path / "a"), (tmp_path / "b").as_uri()):
        writer = ParquetSink().open_table(uri, "t", roll_rows=1)
        writer.write_all(_batches(1, 2))
        assert writer.close() == 2
    assert len(os.listdir(tmp_path / "a" / "t")) == 2


# ---- #742: a damaged workbook is a WorkbookError that names the file ---------------------------


def _damaged_books(tmp_path: Path) -> dict[str, Path]:
    import zipfile

    openpyxl = pytest.importorskip("openpyxl")
    wb = openpyxl.Workbook()
    wb.active.append(["a"])
    wb.active.append([1])
    good = tmp_path / "ok.xlsx"
    wb.save(good)

    def mutate(name: str, part: str, change) -> Path:  # type: ignore[no-untyped-def]
        out = tmp_path / name
        with zipfile.ZipFile(good) as src, zipfile.ZipFile(out, "w") as dst:
            for info in src.infolist():
                data = src.read(info.filename)
                dst.writestr(info, change(data) if info.filename == part else data)
        return out

    books = {
        "cut sheet": mutate("cut_sheet.xlsx", "xl/worksheets/sheet1.xml", lambda b: b[:100]),
        "cut workbook": mutate("cut_book.xlsx", "xl/workbook.xml", lambda b: b[:60]),
        "empty sheet": mutate("empty.xlsx", "xl/worksheets/sheet1.xml", lambda b: b""),
        "junk types": mutate("types.xlsx", "[Content_Types].xml", lambda b: b"junk"),
    }
    raw = bytearray(good.read_bytes())
    with zipfile.ZipFile(good) as zf:
        info = zf.getinfo("xl/worksheets/sheet1.xml")
    offset = info.header_offset + 30 + len(info.filename.encode()) + len(info.extra) + 5
    raw[offset : offset + 6] = b"\xff\xfe\xfd\xfc\xfb\xfa"
    flipped = tmp_path / "flip.xlsx"
    flipped.write_bytes(bytes(raw))
    books["flipped bytes"] = flipped
    return books


def test_a_damaged_workbook_is_a_workbook_error_naming_the_file(tmp_path: Path) -> None:
    from shape.io import WorkbookError, open_workbook, read_table
    from shape.io.excel import read_workbook, workbook_sheet_names

    for label, path in _damaged_books(tmp_path).items():
        for call in (
            lambda p=path: read_table(p),
            lambda p=path: read_table(f"{p}#nosuch"),
            lambda p=path: read_workbook(p),
            lambda p=path: workbook_sheet_names(p),
            lambda p=path: open_workbook(p),
        ):
            with pytest.raises(WorkbookError, match=rf"{path.name}.*damaged"):
                call()
        assert label  # the mutation that made it


def test_an_intact_workbook_still_reads_after_the_damaged_ones(tmp_path: Path) -> None:
    from shape.io import read_table

    _damaged_books(tmp_path)
    assert read_table(tmp_path / "ok.xlsx").to_pydict() == {"a": [1]}


# ---- #743: an empty path is not the current directory -----------------------------------------


@pytest.mark.parametrize("text", ["", "   "])
def test_an_empty_or_blank_path_is_a_missing_source(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, text: str
) -> None:
    from shape.io import expand_paths, open_source, read_table

    pq.write_table(pa.table({"a": [1]}), tmp_path / "p.parquet")
    monkeypatch.chdir(tmp_path)
    for call in (
        lambda: expand_paths(text),
        lambda: expand_paths(["p.parquet", text]),
        lambda: read_table(text),
        lambda: open_source(text),
    ):
        with pytest.raises(FileNotFoundError, match="source not found"):
            call()


def test_dot_still_names_the_current_directory(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    from shape.io import expand_paths, read_table

    pq.write_table(pa.table({"a": [1]}), tmp_path / "p.parquet")
    monkeypatch.chdir(tmp_path)
    assert [p.name for p in expand_paths(".")] == ["p.parquet"]
    assert [p.name for p in expand_paths(Path("."))] == ["p.parquet"]
    assert read_table(".").num_rows == 1


# ---- #745: an empty output location is an error -----------------------------------------------


def test_sinks_refuse_an_empty_location(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    from shape.builtins import sinks

    monkeypatch.chdir(tmp_path)
    calls = {
        "csv": lambda: sinks.CsvSink().write("", "t", _batches(1)),
        "jsonl": lambda: sinks.JsonlSink().write("", "t", _batches(1)),
        "parquet": lambda: sinks.ParquetSink().write("", "t", _batches(1)),
        "rolling": lambda: sinks.ParquetSink().write("", "t", _batches(1), roll_rows=1),
        "open_table": lambda: sinks.CsvSink().open_table("", "t", roll_rows=1),
        "sql": lambda: sinks.SqlSink().write("", "t", _batches(1)),
        "excel": lambda: sinks.ExcelSink().write("", "t", _batches(1)),
        "workbook": lambda: sinks.ExcelSink().write_workbook("", {"t": pa.table({"a": [1]})}),
        "mirror": lambda: sinks.FabricMirrorSink().write("", "t", _batches(1)),
    }
    if __import__("importlib").util.find_spec("deltalake"):
        calls["delta"] = lambda: sinks.DeltaSink().write("", "t", _batches(1))
    for name, call in calls.items():
        with pytest.raises(ValueError, match="output location is empty"):
            call()
        assert os.listdir(tmp_path) == [], name


def test_sinks_still_write_to_dot_and_to_a_directory(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    from shape.builtins import sinks

    monkeypatch.chdir(tmp_path)
    assert sinks.CsvSink().write(".", "t", _batches(1)) == 1
    assert sinks.CsvSink().write("./", "u", _batches(1)) == 1
    assert sorted(os.listdir(tmp_path)) == ["t.csv", "u.csv"]


_ = dt


# ---- #729: overwrite replaces the parts of an earlier, larger rolling run -------------------


def test_rolling_overwrite_removes_stale_parts(tmp_path: Path) -> None:
    sink = ParquetSink()
    sink.write(str(tmp_path), "t", _batches(1, 2, 3, 4, 5, 6), roll_rows=2)
    (tmp_path / "t" / "_SUCCESS").write_bytes(b"")
    (tmp_path / "t" / "notes.txt").write_text("keep")
    (tmp_path / "other").mkdir()
    sink.write(str(tmp_path), "other", _batches(7), roll_rows=2)
    sink.write(str(tmp_path), "t", _batches(9, 9), roll_rows=2)
    assert sorted(os.listdir(tmp_path / "t")) == ["_SUCCESS", "notes.txt", "t-00001.parquet"]
    assert pq.read_table(tmp_path / "t" / "t-00001.parquet").column("a").to_pylist() == [9, 9]
    assert os.listdir(tmp_path / "other") == ["other-00001.parquet"]


def test_rolling_overwrite_keeps_old_parts_when_the_run_fails(tmp_path: Path) -> None:
    sink = ParquetSink()
    sink.write(str(tmp_path), "t", _batches(1, 2, 3, 4, 5, 6), roll_rows=2)

    def boom() -> Any:
        yield _batches(9, 9)[0]
        raise RuntimeError("source failed")

    with pytest.raises(RuntimeError):
        sink.write(str(tmp_path), "t", boom(), roll_rows=2)
    assert len(os.listdir(tmp_path / "t")) == 3


def test_rolling_append_and_fail_keep_existing_parts(tmp_path: Path) -> None:
    sink = ParquetSink()
    sink.write(str(tmp_path), "t", _batches(1, 2, 3, 4), roll_rows=2)
    sink.write(str(tmp_path), "t", _batches(5, 6), roll_rows=2, mode="append")
    assert len(os.listdir(tmp_path / "t")) == 3


# ---- #746: a plain file sink honours mode; the delimiter error is clear ----------------------


def test_plain_file_sink_modes(tmp_path: Path) -> None:
    target = str(tmp_path / "x.csv")
    CsvSink().write(target, "x", _batches(1))
    with pytest.raises(FileExistsError, match="mode=fail"):
        CsvSink().write(target, "x", _batches(2), mode="fail")
    assert "1" in (tmp_path / "x.csv").read_text()
    with pytest.raises(ValueError, match="mode=append needs rolling files"):
        CsvSink().write(target, "x", _batches(3), mode="append")
    with pytest.raises(ValueError, match="mode is one of overwrite, append, fail, got 'bogus'"):
        CsvSink().write(target, "x", _batches(4), mode="bogus")
    CsvSink().write(str(tmp_path / "new.csv"), "x", _batches(5), mode="fail")
    assert (tmp_path / "new.csv").exists()
    CsvSink().write(target, "x", _batches(6), mode="overwrite")
    assert "6" in (tmp_path / "x.csv").read_text()


def test_csv_delimiter_must_be_one_character(tmp_path: Path) -> None:
    with pytest.raises(ValueError, match="the CSV delimiter must be one character, got ';;'"):
        CsvSink().write(str(tmp_path / "y.csv"), "y", _batches(1), delimiter=";;")
    assert not (tmp_path / "y.csv").exists()
