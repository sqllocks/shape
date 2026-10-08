"""P4-06: writers. Every format is a ``shape.sinks`` plugin; golden output per format."""

from __future__ import annotations

import json
import re
import subprocess
import sys
from pathlib import Path

import pyarrow as pa
import pyarrow.parquet as pq
import pytest
import sqlglot
from engine_fixtures import STRATEGIES
from gen_fixtures import schema
from writers_fixtures import COLUMNS, table

from shape.builtins.sinks import SqlSink
from shape.generation.engine import Engine
from shape.generation.output import (
    FORMATS,
    format_summary,
    needs_post_pass,
    write_engine,
    write_result,
)
from shape.plugins.host import default_host

sys.path.insert(0, str(Path(__file__).resolve().parents[2] / "scripts"))
from refengine_name import names_refengine  # noqa: E402

GOLDEN = Path(__file__).parent / "golden" / "writers"
DIALECTS = ("tsql", "tsql-fabric-warehouse", "postgres", "mysql")
SQLGLOT = {
    "tsql": "tsql",
    "tsql-fabric-warehouse": "tsql",
    "postgres": "postgres",
    "mysql": "mysql",
}


def _write(sink_name: str, tmp_path: Path, name: str = "t.out", **options) -> Path:
    sink = default_host().get("shape.sinks", sink_name)
    target = tmp_path / name
    t = table()
    assert sink.write(str(target), "item", iter(t.to_batches(max_chunksize=2)), **options) == 3
    return target


def _golden(name: str, actual: str) -> None:
    expected = (GOLDEN / name).read_text(encoding="utf-8")
    assert actual == expected, f"{name} differs from its golden file"


# ---- registry ------------------------------------------------------------------------------


def test_every_core_and_plugin_format_is_a_registered_sink():
    names = set(default_host().names("shape.sinks"))
    assert {"csv", "tsv", "jsonl", "parquet", "sql", "excel", "delta"} <= names
    for fmt in FORMATS:
        assert fmt == "summary" or fmt in names


# ---- golden output -------------------------------------------------------------------------


@pytest.mark.parametrize("fmt", ["csv", "tsv", "jsonl"])
def test_text_formats_match_golden(fmt, tmp_path):
    path = _write(fmt, tmp_path, f"item.{fmt}")
    _golden(f"item.{fmt}", path.read_text(encoding="utf-8"))


@pytest.mark.parametrize("dialect", DIALECTS)
def test_sql_matches_golden(dialect, tmp_path):
    path = _write(
        "sql",
        tmp_path,
        "item.sql",
        sql_dialect=dialect,
        columns=COLUMNS,
        primary_key=["id"],
        schema_name="dbo",
        batch_size=2,
        header=["Domain: demo | Mode: 3nf | Seed: 7"],
    )
    _golden(f"item.{dialect}.sql", path.read_text(encoding="utf-8"))


def test_summary_matches_golden():
    from shape.generation.engine import GenerationResult

    s = schema()
    result = Engine(s, strategies=STRATEGIES).generate()
    fixed = GenerationResult(
        tables=result.tables,
        schema=result.schema,
        generation_order=result.generation_order,
        elapsed_seconds=0.0,
        row_counts=result.row_counts,
    )
    _golden("summary.txt", format_summary(fixed) + "\n")


# ---- round trips ---------------------------------------------------------------------------


def test_csv_and_tsv_read_back(tmp_path):
    import pyarrow.csv as pacsv

    for fmt, delim in (("csv", ","), ("tsv", "\t")):
        path = _write(fmt, tmp_path, f"item.{fmt}")
        back = pacsv.read_csv(
            str(path),
            parse_options=pacsv.ParseOptions(delimiter=delim),
            convert_options=pacsv.ConvertOptions(strings_can_be_null=True),
        )
        assert back.column_names == table().column_names
        assert back["id"].to_pylist() == [1, 2, 3]
        assert back["name"].to_pylist() == ["plain", "O'Brien \\ é", None]


def test_jsonl_is_one_object_per_line_with_null(tmp_path):
    rows = [
        json.loads(x) for x in _write("jsonl", tmp_path, "i.jsonl").read_text("utf-8").splitlines()
    ]
    assert [r["id"] for r in rows] == [1, 2, 3] and rows[2]["name"] is None


def test_parquet_reads_back_in_pandas_with_t17_settings(tmp_path):
    import pandas as pd

    path = _write("parquet", tmp_path, "item.parquet")
    frame = pd.read_parquet(path)
    assert list(frame.columns) == table().column_names and len(frame) == 3
    assert pq.read_table(path).equals(table())
    meta = pq.ParquetFile(path).metadata.row_group(0).column(0)
    assert meta.compression == "SNAPPY"
    assert any("DICTIONARY" in str(e) for e in meta.encodings)


def test_excel_reads_back(tmp_path):
    import openpyxl

    path = _write("excel", tmp_path, "item.xlsx")
    book = openpyxl.load_workbook(path, read_only=True)
    try:
        rows = list(book["item"].iter_rows(values_only=True))
    finally:
        book.close()
    assert rows[0] == tuple(table().column_names) and len(rows) == 4
    assert rows[1][1] == "plain" and rows[3][1] is None


def test_delta_reads_back(tmp_path):
    import deltalake

    sink = default_host().get("shape.sinks", "delta")
    t = table()
    assert sink.write(str(tmp_path), "item", iter(t.to_batches(max_chunksize=2))) == 3
    back = deltalake.DeltaTable(str(tmp_path / "item")).to_pyarrow_table()
    assert back.sort_by("id")["id"].to_pylist() == [1, 2, 3]


def test_empty_batches_still_give_a_valid_file(tmp_path):
    sink = default_host().get("shape.sinks", "csv")
    sink.write(str(tmp_path / "e.csv"), "e", iter([]), schema=table().schema)
    assert (tmp_path / "e.csv").exists()


# ---- SQL ----------------------------------------------------------------------------------


def _statements(text: str) -> list[str]:
    body = "\n".join(line for line in text.splitlines() if not line.startswith("--"))
    return [p.strip() for p in re.split(r"^GO\s*$", body, flags=re.M) if p.strip()]


@pytest.mark.parametrize("dialect", DIALECTS)
def test_sql_parses_with_sqlglot_for_each_dialect(dialect, tmp_path):
    path = _write(
        "sql",
        tmp_path,
        "item.sql",
        sql_dialect=dialect,
        columns=COLUMNS,
        primary_key=["id"],
        schema_name="dbo",
        batch_size=2,
    )
    text = path.read_text(encoding="utf-8")
    chunks = _statements(text)
    parsed = []
    for chunk in chunks:
        parsed += sqlglot.parse(chunk, read=SQLGLOT[dialect], error_level=sqlglot.ErrorLevel.RAISE)
    kinds = [type(p).__name__ for p in parsed if p is not None]
    assert kinds.count("Create") == 1 and kinds.count("Insert") == 2
    assert "Drop" in kinds or any("DROP" in c.upper() for c in chunks)
    inserted = sum(len(p.expression.expressions) for p in parsed if type(p).__name__ == "Insert")
    assert inserted == 3


@pytest.mark.parametrize("dialect", DIALECTS)
def test_sql_names_and_headers_cannot_add_statements(dialect, tmp_path):
    # A line break in a name or header line would end a -- comment; a quote would end a literal.
    hostile = "it'em]`\"\nDELETE FROM z; --\r\nDROP TABLE y; --"
    target = tmp_path / "hostile.sql"
    t = table()
    write = lambda: SqlSink().write(  # noqa: E731
        str(target),
        hostile,
        iter(t.to_batches()),
        sql_dialect=dialect,
        columns=COLUMNS,
        primary_key=["id"],
        header=["first\nDELETE FROM z;", "second\rDROP TABLE y;"],
    )
    if dialect.startswith("tsql"):  # #724: a T-SQL script cannot hold a line break in a name
        with pytest.raises(ValueError, match="contains a line break"):
            write()
        assert not target.exists()
        return
    write()
    text = target.read_text(encoding="utf-8")
    parsed = []
    for chunk in _statements(text):
        parsed += sqlglot.parse(chunk, read=SQLGLOT[dialect], error_level=sqlglot.ErrorLevel.RAISE)
    kinds = [type(p).__name__ for p in parsed if p is not None]
    assert "Delete" not in kinds
    assert kinds.count("Create") == 1 and kinds.count("Insert") == 1
    assert kinds.count("Drop") <= 1
    assert "-- first DELETE FROM z;\n-- second DROP TABLE y;\n" in text


def test_sql_options_ddl_drop_go_and_schema_name(tmp_path):
    base = {"sql_dialect": "tsql", "columns": COLUMNS, "primary_key": ["id"]}
    text = _write("sql", tmp_path, "a.sql", **base, ddl=False).read_text("utf-8")
    assert "CREATE TABLE" not in text and "DROP" not in text and "INSERT INTO [item]" in text
    text = _write("sql", tmp_path, "b.sql", **base, drop=False).read_text("utf-8")
    assert "CREATE TABLE" in text and "DROP" not in text and "OBJECT_ID" not in text
    text = _write("sql", tmp_path, "c.sql", **base, go=False).read_text("utf-8")
    assert "\nGO" not in text
    text = _write("sql", tmp_path, "d.sql", **base, schema_name="sales").read_text("utf-8")
    assert "[sales].[item]" in text
    text = _write("sql", tmp_path, "e.sql", sql_dialect="postgres", go=True).read_text("utf-8")
    assert "GO" not in text  # GO is T-SQL only


def test_sql_without_schema_metadata_types_from_arrow(tmp_path):
    text = _write("sql", tmp_path, "m.sql", sql_dialect="postgres").read_text("utf-8")
    assert '"id"' in text and "BIGINT" in text and "NUMERIC(10,2)" in text and "TIMESTAMP" in text


def test_tsql_batches_never_exceed_1000_rows(tmp_path):
    sink = SqlSink()
    big = pa.table({"x": pa.array(range(2500), pa.int64())})
    sink.write(str(tmp_path / "big.sql"), "big", iter(big.to_batches()), batch_size=5000, ddl=False)
    text = (tmp_path / "big.sql").read_text("utf-8")
    assert text.count("INSERT INTO") == 3
    sink.write(
        str(tmp_path / "pg.sql"),
        "big",
        iter(big.to_batches()),
        batch_size=5000,
        ddl=False,
        sql_dialect="postgres",
    )
    assert (tmp_path / "pg.sql").read_text("utf-8").count("INSERT INTO") == 1


def test_sql_rejects_unknown_dialect(tmp_path):
    with pytest.raises(ValueError, match="unknown SQL dialect"):
        _write("sql", tmp_path, "x.sql", sql_dialect="oracle")


def test_sql_is_deterministic_and_names_no_other_product(tmp_path):
    one = _write("sql", tmp_path, "1.sql", columns=COLUMNS).read_bytes()
    two = _write("sql", tmp_path, "2.sql", columns=COLUMNS).read_bytes()
    assert one == two and not names_refengine(one)


def test_sql_nan_and_inf_become_null(tmp_path):
    t = pa.table({"x": pa.array([float("nan"), float("inf"), 1.5])})
    SqlSink().write(str(tmp_path / "n.sql"), "n", iter(t.to_batches()), ddl=False)
    text = (tmp_path / "n.sql").read_text("utf-8")
    assert text.count("NULL") == 2 and "1.5" in text


# ---- through the engine ---------------------------------------------------------------------


def _engine(**kw) -> Engine:
    return Engine(schema(), strategies=STRATEGIES, **kw)


def test_write_result_writes_every_table_in_every_core_format(tmp_path):
    result = _engine().generate()
    for fmt in ("csv", "tsv", "jsonl", "parquet", "sql"):
        out = tmp_path / fmt
        files = write_result(result, fmt, out)
        assert sorted(p.stem for p in files) == sorted(result.tables)
        assert all(p.exists() for p in files)
    back = pq.read_table(tmp_path / "parquet" / "order.parquet")
    assert back.equals(result["order"])


def test_sql_from_a_schema_carries_keys_and_nullability(tmp_path):
    result = _engine().generate()
    write_result(result, "sql", tmp_path, sql_dialect="postgres")
    text = (tmp_path / "customer.sql").read_text("utf-8")
    assert 'PRIMARY KEY ("customer_id")' in text and '"name"' in text
    assert re.search(r'"name"\s+VARCHAR\(255\)\s+NULL', text)  # null_rate 0.2 -> nullable


def test_unknown_format_is_refused(tmp_path):
    with pytest.raises(ValueError, match="unknown format"):
        write_result(_engine().generate(), "xml", tmp_path)


def test_streaming_matches_the_whole_table_and_is_chunk_independent(tmp_path):
    s = schema()
    for t in s.tables.values():  # drop the post-pass columns: this schema streams
        for name in [n for n, c in t.columns.items() if c.strategy == "computed"]:
            del t.columns[name]
    s.business_rules = []
    s.correlated_columns = {}
    assert not needs_post_pass(s)
    outputs = []
    for chunk in (7, 64, 1000):
        out = tmp_path / str(chunk)
        write_engine(Engine(s, strategies=STRATEGIES), "parquet", out, chunk_rows=chunk)
        outputs.append({p.stem: pq.read_table(p) for p in out.glob("*.parquet")})
    assert outputs[0].keys() == outputs[1].keys() == outputs[2].keys()
    for name in outputs[0]:
        assert outputs[0][name].equals(outputs[1][name]) and outputs[1][name].equals(
            outputs[2][name]
        )
    whole = Engine(s, strategies=STRATEGIES).generate()
    for name, got in outputs[0].items():
        assert got.equals(whole[name])


def test_write_engine_falls_back_for_post_passes(tmp_path):
    s = schema()
    assert needs_post_pass(s)  # the fixture schema has a computed column
    write_engine(Engine(s, strategies=STRATEGIES), "parquet", tmp_path)
    expected = Engine(s, strategies=STRATEGIES).generate()
    assert pq.read_table(tmp_path / "order.parquet").equals(expected["order"])


def test_a_failing_producer_surfaces_in_write_engine(tmp_path):
    s = schema()
    for t in s.tables.values():
        for name in [n for n, c in t.columns.items() if c.strategy == "computed"]:
            del t.columns[name]
    s.business_rules = []
    s.correlated_columns = {}
    engine = Engine(s, strategies=STRATEGIES)
    original = engine.generate_chunk

    def boom(table_name, start, n, **kw):
        if start > 0:
            raise RuntimeError("generator failed")
        return original(table_name, start, n, **kw)

    engine.generate_chunk = boom  # type: ignore[method-assign]
    with pytest.raises(RuntimeError, match="generator failed"):
        write_engine(engine, "csv", tmp_path, chunk_rows=10)


def test_written_files_are_closed_so_windows_can_delete_them(tmp_path):
    result = _engine().generate()
    for fmt in ("csv", "parquet", "jsonl", "sql", "tsv"):
        out = tmp_path / fmt
        for p in write_result(result, fmt, out):
            p.rename(p.with_name(p.name + ".moved"))  # fails on Windows while a handle is open
            p.with_name(p.name + ".moved").unlink()


def test_optional_extras_are_imported_lazily():
    code = (
        "import sys, shape.builtins.sinks as s; "
        "assert 'openpyxl' not in sys.modules and 'deltalake' not in sys.modules"
    )
    subprocess.run([sys.executable, "-c", code], check=True)
